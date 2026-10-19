import asyncio
import hashlib
import hmac
import io
import json
import logging
import secrets
import time
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from typing import Any

from fastapi import Depends, FastAPI, File, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, PlainTextResponse
import httpx
import asyncpg
import asyncmy
import jwt
from pypdf import PdfReader
from docx import Document as DocxDocument
from openpyxl import load_workbook
from redis.asyncio import Redis
from sqlalchemy import delete, desc, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from .config import get_settings
from .db import AdminSessionLocal, Base, SessionLocal, admin_engine, engine, get_admin_db, get_db
from .deps import AdminContext, Context, EditorContext, OperatorContext, PermissionAdminContext, RequestContext
from .models import (
    Agent,
    Application,
    ApplicationVersion,
    AuditEvent,
    Document,
    DocumentChunk,
    ComputeQuota,
    ComputeUsage,
    EvaluationJob,
    GpuLease,
    KnowledgeBase,
    KnowledgeBaseAccess,
    MCPServer,
    Membership,
    ModelHealthCheck,
    ModelProvider,
    Plugin,
    Project,
    Role,
    ScimToken,
    Skill,
    Tenant,
    TenantPolicy,
    User,
    Workflow,
    WorkflowRelease,
    WorkflowRun,
    WorkflowSchedule,
    WorkflowVersion,
    SandboxRun,
    KnowledgeAsset,
    AgentEvent,
    ApiKey,
    AgentConfig,
    AgentVersion,
    ApprovalRequest as ApprovalRequestModel,
    DatabaseConnection,
    Favorite,
    InterventionRule,
    MarketplaceItem,
    Organization,
    PromptTemplate,
    PublishBinding,
    PublishChannel,
    Workspace,
    WorkspaceMember,
)
from .schemas import (
    AgentChatRequest,
    AgentCreate,
    ApplicationCreate,
    ApplicationVersionCreate,
    ApplicationRollback,
    ApplicationUpdate,
    ApprovalRequest,
    KnowledgeBaseCreate,
    LDAPLoginRequest,
    LoginRequest,
    LoginResponse,
    MemberCreate,
    MemberRoleUpdate,
    MCPServerCreate,
    ModelTestRequest,
    OIDCCallbackRequest,
    OIDCStartRequest,
    PasswordChangeRequest,
    MFACodeRequest,
    MFADisableRequest,
    ScimTokenCreate,
    ScimUserCreate,
    ScimUserPatch,
    PluginCreate,
    ProjectCreate,
    ProviderCreate,
    QuotaUpdate,
    TenantPolicyUpdate,
    PublicAppInvokeRequest,
    EvaluationRequest,
    EvaluationJobCreate,
    GpuLeaseRequest,
    SandboxRunCreate,
    BudgetUpdate,
    ReleaseRequest,
    RollbackRequest,
    ScheduleCreate,
    ScheduleUpdate,
    SkillCreate,
    WorkflowCreate,
    WorkflowRunRequest,
    WorkflowWebhookRequest,
    WorkflowVersionCreate,
    EnabledUpdate,
    KnowledgeBaseAccessCreate,
    RoleCreate,
    RoleUpdate,
    AgentConfigUpdate,
    AgentPublishRequest,
    ApprovalDecision,
    DatabaseCreate,
    DatabaseQuery,
    FavoriteRequest,
    InterventionRuleCreate,
    MarketplaceCreate,
    ModelExperienceRequest,
    OrganizationCreate,
    PromptTemplateCreate,
    PublishBindingCreate,
    WorkspaceCreate,
    KnowledgeAssetCreate,
    AgentFeedbackCreate,
    ApiKeyCreate,
)
from .security import create_access_token, decrypt_secret, encrypt_secret, hash_password, secret_backend_status, verify_password
from .services.crypto import checksum, chunk_text
from .services.knowledge import search_chunks
from .services.identity import BUILTIN_ROLES, mapped_role, parse_role_mapping
from .services.ldap_auth import LDAPError, authenticate as ldap_authenticate
from .services.llm import call_chat
from .services.oidc import OIDCError, create_authorization_url, exchange_code
from .services.observability import metrics_payload, setup_telemetry
from .services.safety import inspect_text
from .services.vector import embed
from .services.alerts import send_alert
from .services.sandbox import build_job_manifest
from .services.totp import provisioning_uri, random_base32, verify as verify_totp
from .services.workflow import enqueue_index, enqueue_run, ordered_nodes


logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s %(message)s")
logger = logging.getLogger("enterprise_ai_platform")

PERMISSION_CATALOG = [
    {"key": "workspace.read", "name": "查看工作区"},
    {"key": "workflow.write", "name": "编辑工作流"},
    {"key": "workflow.run", "name": "运行工作流"},
    {"key": "knowledge.write", "name": "管理知识库"},
    {"key": "model.admin", "name": "管理模型与算力"},
    {"key": "extension.admin", "name": "管理技能与插件"},
    {"key": "member.admin", "name": "管理成员与权限"},
    {"key": "audit.read", "name": "查看审计与观测"},
]


def utcnow() -> datetime:
    return datetime.now(timezone.utc)


async def audit(db: AsyncSession, ctx: RequestContext | None, action: str, resource_type: str = "", resource_id: str = "", detail: dict[str, Any] | None = None) -> None:
    db.add(AuditEvent(tenant_id=ctx.tenant_id if ctx else "system", actor_id=ctx.user_id if ctx else None, action=action, resource_type=resource_type, resource_id=resource_id, detail_json=detail or {}))


async def tenant_policy(db: AsyncSession, tenant_id: str) -> TenantPolicy:
    row = await db.scalar(select(TenantPolicy).where(TenantPolicy.tenant_id == tenant_id))
    if row:
        return row
    settings = get_settings()
    row = TenantPolicy(
        tenant_id=tenant_id,
        safety_json={
            "mode": settings.safety_mode,
            "max_input_chars": settings.safety_max_input_chars,
            "block_prompt_injection": settings.safety_block_prompt_injection,
        },
        pricing_json={
            "currency": settings.finops_currency,
            "token_price_per_1k": settings.finops_token_price_per_1k,
            "gpu_price_per_second": settings.finops_gpu_price_per_second,
            "cpu_price_per_second": settings.finops_cpu_price_per_second,
        },
    )
    db.add(row)
    await db.flush()
    return row


def validate_extension_manifest(manifest: dict[str, Any]) -> None:
    capabilities = manifest.get("capabilities") or []
    if not isinstance(capabilities, list):
        raise HTTPException(status_code=422, detail="扩展 capabilities 必须是数组")
    invalid = sorted(set(str(item) for item in capabilities) - get_settings().extension_capability_list)
    if invalid:
        raise HTTPException(status_code=422, detail=f"扩展能力未获批准: {', '.join(invalid)}")


async def bootstrap() -> None:
    settings = get_settings()
    async with admin_engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    async with AdminSessionLocal() as db:
        tenant = await db.scalar(select(Tenant).where(Tenant.slug == settings.bootstrap_tenant_slug))
        if not tenant:
            tenant = Tenant(name=settings.bootstrap_tenant, slug=settings.bootstrap_tenant_slug)
            db.add(tenant)
            await db.flush()
        user = await db.scalar(select(User).where(User.email == settings.bootstrap_admin_email))
        if not user:
            user = User(email=settings.bootstrap_admin_email, display_name="Platform Administrator", password_hash=hash_password(settings.bootstrap_admin_password))
            db.add(user)
            await db.flush()
        membership = await db.scalar(select(Membership).where(Membership.user_id == user.id, Membership.tenant_id == tenant.id))
        if not membership:
            db.add(Membership(user_id=user.id, tenant_id=tenant.id, role="owner"))
        project = await db.scalar(select(Project).where(Project.tenant_id == tenant.id, Project.slug == "platform"))
        if not project:
            db.add(Project(tenant_id=tenant.id, name="Platform Core", slug="platform", description="企业 AI 中台基础项目"))
        organization = await db.scalar(select(Organization).where(Organization.tenant_id == tenant.id, Organization.parent_id.is_(None), Organization.name == "默认组织"))
        if not organization:
            organization = Organization(tenant_id=tenant.id, name="默认组织", code="default")
            db.add(organization)
            await db.flush()
        workspace = await db.scalar(select(Workspace).where(Workspace.tenant_id == tenant.id, Workspace.slug == "default"))
        if not workspace:
            workspace = Workspace(tenant_id=tenant.id, organization_id=organization.id, owner_id=user.id, name="默认工作区", slug="default", workspace_type="team", description="企业 AI 中台默认工作区")
            db.add(workspace)
            await db.flush()
        workspace_member = await db.scalar(select(WorkspaceMember).where(WorkspaceMember.workspace_id == workspace.id, WorkspaceMember.user_id == user.id))
        if not workspace_member:
            db.add(WorkspaceMember(tenant_id=tenant.id, workspace_id=workspace.id, user_id=user.id, role="owner"))
        for code, name, kind in [
            ("api", "Open API", "api"), ("mcp", "MCP 服务", "mcp"), ("web", "Web 应用", "web"),
            ("websdk", "WebSDK", "websdk"), ("feishu", "飞书", "im"), ("wecom", "企业微信", "im"),
            ("wechat", "微信公众号", "im"), ("a2a", "A2A 服务", "a2a"),
        ]:
            channel = await db.scalar(select(PublishChannel).where(PublishChannel.tenant_id == tenant.id, PublishChannel.code == code))
            if not channel:
                db.add(PublishChannel(tenant_id=tenant.id, name=name, code=code, channel_type=kind, approval_required=kind in {"im", "a2a"}, created_by=user.id))
        built_in_roles = {
            "owner": [item["key"] for item in PERMISSION_CATALOG],
            "admin": [item["key"] for item in PERMISSION_CATALOG if item["key"] != "member.admin"],
            "operator": ["workspace.read", "workflow.run", "audit.read"],
            "developer": ["workspace.read", "workflow.write", "workflow.run", "knowledge.write"],
            "viewer": ["workspace.read", "audit.read"],
        }
        for slug, permissions in built_in_roles.items():
            role = await db.scalar(select(Role).where(Role.tenant_id == tenant.id, Role.slug == slug))
            if not role:
                db.add(Role(tenant_id=tenant.id, name={"owner": "所有者", "admin": "管理员", "operator": "操作员", "developer": "开发者", "viewer": "查看者"}[slug], slug=slug, permissions_json=permissions, is_system=True, created_by=user.id))
        await db.commit()


@asynccontextmanager
async def lifespan(app: FastAPI):
    setup_telemetry()
    await bootstrap()
    yield
    await engine.dispose()
    await admin_engine.dispose()


class UTF8JSONResponse(JSONResponse):
    media_type = "application/json; charset=utf-8"


app = FastAPI(title="Enterprise AI Platform", version="0.1.0", lifespan=lifespan, default_response_class=UTF8JSONResponse)
app.add_middleware(CORSMiddleware, allow_origins=get_settings().cors_origin_list, allow_credentials=False, allow_methods=["*"], allow_headers=["*"])


@app.middleware("http")
async def security_middleware(request, call_next):
    """Apply baseline browser security headers and Redis-backed API throttling."""
    if request.url.path.startswith("/api/"):
        window = 60
        limit = 20 if request.url.path.startswith("/api/auth/") else 300
        client_ip = request.client.host if request.client else "unknown"
        redis = Redis.from_url(get_settings().redis_url, decode_responses=True)
        try:
            bucket = f"ai_platform:rate:{client_ip}:{request.url.path.split('/')[2]}"
            count = await redis.incr(bucket)
            if count == 1:
                await redis.expire(bucket, window)
            if count > limit:
                response = JSONResponse({"detail": "请求过于频繁，请稍后重试"}, status_code=429)
                response.headers["Retry-After"] = str(window)
                return response
        except Exception:
            logger.warning("rate limiter unavailable", exc_info=True)
        finally:
            await redis.aclose()
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    if get_settings().app_env == "production":
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    return response


@app.get("/healthz")
async def healthz(db: AsyncSession = Depends(get_db)):
    await db.execute(select(1))
    return {"ok": True, "service": "enterprise-ai-platform", "time": utcnow().isoformat()}


@app.get("/readyz")
async def readyz(db: AsyncSession = Depends(get_db)):
    await db.execute(select(1))
    redis = Redis.from_url(get_settings().redis_url, decode_responses=True)
    try:
        await redis.ping()
    finally:
        await redis.aclose()
    return {"ok": True, "service": "enterprise-ai-platform", "dependencies": {"postgres": "ok", "redis": "ok"}, "time": utcnow().isoformat()}


@app.get("/metrics", response_class=PlainTextResponse)
async def metrics():
    return PlainTextResponse(metrics_payload().decode(), media_type="text/plain; version=0.0.4")


@app.post("/api/auth/login", response_model=LoginResponse)
async def login(body: LoginRequest, db: AsyncSession = Depends(get_admin_db)):
    user = await db.scalar(select(User).where(User.email == body.email, User.is_active.is_(True)))
    if not user or not verify_password(body.password, user.password_hash):
        raise HTTPException(status_code=401, detail="邮箱或密码错误")
    if user.mfa_enabled:
        if not body.mfa_code:
            raise HTTPException(status_code=401, detail={"code": "mfa_required", "message": "请输入多因素验证码"})
        try:
            valid = verify_totp(decrypt_secret(user.mfa_secret_encrypted), body.mfa_code, window=1)
        except Exception:
            valid = False
        if not valid:
            raise HTTPException(status_code=401, detail={"code": "mfa_invalid", "message": "多因素验证码无效"})
    stmt = select(Membership, Tenant).join(Tenant, Tenant.id == Membership.tenant_id).where(Membership.user_id == user.id, Tenant.is_active.is_(True))
    if body.tenant_slug:
        stmt = stmt.where(Tenant.slug == body.tenant_slug)
    row = (await db.execute(stmt)).first()
    if not row:
        raise HTTPException(status_code=403, detail="用户没有可访问的租户")
    membership, tenant = row
    token = create_access_token(user.id, tenant.id, membership.role)
    return LoginResponse(access_token=token, user={"id": user.id, "email": user.email, "display_name": user.display_name, "role": membership.role}, tenant={"id": tenant.id, "name": tenant.name, "slug": tenant.slug})


@app.post("/api/auth/oidc/start")
async def oidc_start(body: OIDCStartRequest):
    settings = get_settings()
    tenant_slug = body.tenant_slug or settings.bootstrap_tenant_slug
    try:
        authorization_url, state_payload = await create_authorization_url(tenant_slug)
    except (OIDCError, httpx.HTTPError) as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    state = authorization_url.split("state=", 1)[1].split("&", 1)[0]
    redis = Redis.from_url(settings.redis_url, decode_responses=True)
    try:
        await redis.setex(f"oidc:state:{state}", 600, state_payload)
    finally:
        await redis.aclose()
    return {"authorization_url": authorization_url, "state": state, "expires_in": 600}


@app.post("/api/auth/ldap/login", response_model=LoginResponse)
async def ldap_login(body: LDAPLoginRequest, db: AsyncSession = Depends(get_admin_db)):
    settings = get_settings()
    try:
        identity = await asyncio.to_thread(ldap_authenticate, body.email, body.password)
    except LDAPError as exc:
        raise HTTPException(status_code=503 if "尚未配置" in str(exc) else 401, detail=str(exc)) from exc
    tenant_slug = body.tenant_slug or settings.bootstrap_tenant_slug
    tenant = await db.scalar(select(Tenant).where(Tenant.slug == tenant_slug, Tenant.is_active.is_(True)))
    if not tenant:
        raise HTTPException(status_code=404, detail="租户不存在")
    email = identity["email"]
    user = await db.scalar(select(User).where(User.email == email))
    if user and not user.is_active:
        raise HTTPException(status_code=403, detail="用户已停用")
    if not user:
        if not settings.ldap_auto_provision:
            raise HTTPException(status_code=403, detail="LDAP 用户尚未加入平台")
        user = User(email=email, display_name=identity["display_name"], password_hash=hash_password(secrets.token_urlsafe(32)))
        db.add(user)
        await db.flush()
    membership = await db.scalar(select(Membership).where(Membership.user_id == user.id, Membership.tenant_id == tenant.id))
    if not membership:
        if not settings.ldap_auto_provision:
            raise HTTPException(status_code=403, detail="LDAP 用户尚未加入该租户")
        membership = Membership(user_id=user.id, tenant_id=tenant.id, role="viewer")
        db.add(membership)
        await db.flush()
    await db.commit()
    token = create_access_token(user.id, tenant.id, membership.role)
    return LoginResponse(access_token=token, user={"id": user.id, "email": user.email, "display_name": user.display_name, "role": membership.role}, tenant={"id": tenant.id, "name": tenant.name, "slug": tenant.slug})


@app.post("/api/auth/oidc/callback", response_model=LoginResponse)
async def oidc_callback(body: OIDCCallbackRequest, db: AsyncSession = Depends(get_admin_db)):
    settings = get_settings()
    redis = Redis.from_url(settings.redis_url, decode_responses=True)
    try:
        state_payload = await redis.get(f"oidc:state:{body.state}")
        await redis.delete(f"oidc:state:{body.state}")
    finally:
        await redis.aclose()
    if not state_payload:
        raise HTTPException(status_code=400, detail="OIDC 状态无效或已过期")
    try:
        state_data = json.loads(state_payload)
        claims = await exchange_code(body.code, state_data["nonce"])
    except (OIDCError, httpx.HTTPError, KeyError, ValueError, jwt.PyJWTError) as exc:
        raise HTTPException(status_code=401, detail=f"OIDC 登录失败: {exc}") from exc

    email_claim = settings.oidc_email_claim
    email = str(claims.get(email_claim) or "").lower().strip()
    if not email:
        raise HTTPException(status_code=401, detail=f"OIDC 身份缺少 {email_claim}")
    if settings.oidc_require_email_verified and claims.get("email_verified") is not True:
        raise HTTPException(status_code=403, detail="OIDC 邮箱尚未验证")
    domain = email.rsplit("@", 1)[-1]
    if settings.oidc_allowed_domain_list and domain not in settings.oidc_allowed_domain_list:
        raise HTTPException(status_code=403, detail="OIDC 邮箱域名不在允许范围内")
    display_name = str(claims.get("name") or claims.get("preferred_username") or email)
    tenant = await db.scalar(select(Tenant).where(Tenant.slug == state_data["tenant_slug"], Tenant.is_active.is_(True)))
    if not tenant:
        raise HTTPException(status_code=404, detail="租户不存在")
    target_role = mapped_role(
        claims,
        settings.oidc_groups_claim,
        settings.oidc_roles_claim,
        parse_role_mapping(settings.oidc_group_role_map),
        settings.oidc_default_role,
    )
    if target_role not in BUILTIN_ROLES:
        custom_role = await db.scalar(select(Role).where(Role.tenant_id == tenant.id, Role.slug == target_role))
        if not custom_role:
            raise HTTPException(status_code=503, detail=f"OIDC 映射角色不存在: {target_role}")
    else:
        custom_role = None
    user = await db.scalar(select(User).where(User.email == email))
    if user and not user.is_active:
        raise HTTPException(status_code=403, detail="用户已停用")
    if not user:
        if not settings.oidc_auto_provision:
            raise HTTPException(status_code=403, detail="用户尚未加入平台")
        user = User(email=email, display_name=display_name, password_hash=hash_password(secrets.token_urlsafe(32)))
        db.add(user)
        await db.flush()
    membership = await db.scalar(select(Membership).where(Membership.user_id == user.id, Membership.tenant_id == tenant.id))
    if not membership:
        if not settings.oidc_auto_provision:
            raise HTTPException(status_code=403, detail="用户尚未加入该租户")
        membership = Membership(
            user_id=user.id,
            tenant_id=tenant.id,
            role=target_role if target_role in BUILTIN_ROLES else "viewer",
            custom_role_id=custom_role.id if custom_role else None,
        )
        db.add(membership)
        await db.flush()
    elif settings.oidc_sync_roles:
        membership.role = target_role if target_role in BUILTIN_ROLES else "viewer"
        membership.custom_role_id = custom_role.id if custom_role else None
    await db.commit()
    token = create_access_token(user.id, tenant.id, membership.role)
    return LoginResponse(access_token=token, user={"id": user.id, "email": user.email, "display_name": user.display_name, "role": membership.role}, tenant={"id": tenant.id, "name": tenant.name, "slug": tenant.slug})


@app.get("/api/me")
async def me(ctx: Context):
    return {"user": {"id": ctx.user.id, "email": ctx.user.email, "display_name": ctx.user.display_name, "mfa_enabled": ctx.user.mfa_enabled}, "tenant": {"id": ctx.tenant.id, "name": ctx.tenant.name, "slug": ctx.tenant.slug}, "role": ctx.role}


@app.post("/api/me/password")
async def change_password(body: PasswordChangeRequest, ctx: Context, db: AsyncSession = Depends(get_db)):
    if not verify_password(body.current_password, ctx.user.password_hash):
        raise HTTPException(status_code=401, detail="当前密码不正确")
    ctx.user.password_hash = hash_password(body.new_password)
    await audit(db, ctx, "user.password.change", "user", ctx.user_id)
    await db.commit()
    return {"changed": True}


@app.get("/api/me/api-keys")
async def list_api_keys(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(ApiKey).where(ApiKey.tenant_id == ctx.tenant_id, ApiKey.user_id == ctx.user_id).order_by(desc(ApiKey.created_at)))).scalars().all()
    return [{"id": row.id, "name": row.name, "prefix": row.key_prefix, "scopes": row.scopes_json, "enabled": row.enabled, "expires_at": row.expires_at, "last_used_at": row.last_used_at, "created_at": row.created_at} for row in rows]


@app.post("/api/me/api-keys")
async def create_api_key(body: ApiKeyCreate, ctx: Context, db: AsyncSession = Depends(get_db)):
    raw = f"eai_{secrets.token_urlsafe(30)}"
    row = ApiKey(tenant_id=ctx.tenant_id, user_id=ctx.user_id, name=body.name, key_prefix=raw[:12], key_hash=hashlib.sha256(raw.encode()).hexdigest(), scopes_json=body.scopes, expires_at=body.expires_at)
    db.add(row)
    try:
        await db.flush()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(status_code=409, detail="API Key 名称已存在") from exc
    await audit(db, ctx, "api_key.create", "api_key", row.id, {"scopes": body.scopes})
    await db.commit()
    return {"id": row.id, "name": row.name, "key": raw, "warning": "完整密钥只在本次响应显示", "scopes": row.scopes_json, "expires_at": row.expires_at}


@app.delete("/api/me/api-keys/{key_id}")
async def revoke_api_key(key_id: str, ctx: Context, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(ApiKey).where(ApiKey.id == key_id, ApiKey.tenant_id == ctx.tenant_id, ApiKey.user_id == ctx.user_id))
    if not row:
        raise HTTPException(status_code=404, detail="API Key 不存在")
    row.enabled = False
    await audit(db, ctx, "api_key.revoke", "api_key", row.id)
    await db.commit()
    return {"id": key_id, "enabled": False}


@app.post("/api/me/mfa/setup")
async def setup_mfa(ctx: Context, db: AsyncSession = Depends(get_db)):
    secret = random_base32()
    ctx.user.mfa_secret_encrypted = encrypt_secret(secret)
    ctx.user.mfa_enabled = False
    await audit(db, ctx, "user.mfa.setup", "user", ctx.user_id)
    await db.commit()
    uri = provisioning_uri(secret, ctx.user.email)
    return {"secret": secret, "otpauth_url": uri, "enabled": False}


@app.post("/api/me/mfa/enable")
async def enable_mfa(body: MFACodeRequest, ctx: Context, db: AsyncSession = Depends(get_db)):
    if not ctx.user.mfa_secret_encrypted:
        raise HTTPException(status_code=422, detail="请先生成 MFA 密钥")
    try:
        valid = verify_totp(decrypt_secret(ctx.user.mfa_secret_encrypted), body.code, window=1)
    except Exception:
        valid = False
    if not valid:
        raise HTTPException(status_code=422, detail="MFA 验证码无效")
    ctx.user.mfa_enabled = True
    await audit(db, ctx, "user.mfa.enable", "user", ctx.user_id)
    await db.commit()
    return {"enabled": True}


@app.post("/api/me/mfa/disable")
async def disable_mfa(body: MFADisableRequest, ctx: Context, db: AsyncSession = Depends(get_db)):
    if not verify_password(body.current_password, ctx.user.password_hash):
        raise HTTPException(status_code=401, detail="当前密码不正确")
    if not ctx.user.mfa_secret_encrypted:
        return {"enabled": False}
    try:
        valid = verify_totp(decrypt_secret(ctx.user.mfa_secret_encrypted), body.code, window=1)
    except Exception:
        valid = False
    if not valid:
        raise HTTPException(status_code=422, detail="MFA 验证码无效")
    ctx.user.mfa_enabled = False
    ctx.user.mfa_secret_encrypted = ""
    await audit(db, ctx, "user.mfa.disable", "user", ctx.user_id)
    await db.commit()
    return {"enabled": False}


@app.get("/api/admin/security/status")
async def security_status(ctx: AdminContext):
    return secret_backend_status()


@app.get("/api/admin/policy")
async def get_policy(ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = await tenant_policy(db, ctx.tenant_id)
    await db.commit()
    return {"tenant_id": ctx.tenant_id, "safety": row.safety_json, "pricing": row.pricing_json}


@app.patch("/api/admin/policy")
async def update_policy(body: TenantPolicyUpdate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = await tenant_policy(db, ctx.tenant_id)
    if body.safety:
        allowed_safety = {"mode", "max_input_chars", "block_prompt_injection"}
        invalid = sorted(set(body.safety) - allowed_safety)
        if invalid:
            raise HTTPException(status_code=422, detail=f"未知安全策略: {', '.join(invalid)}")
        if body.safety.get("mode") not in {None, "off", "redact", "block"}:
            raise HTTPException(status_code=422, detail="安全模式只能是 off、redact 或 block")
        if body.safety.get("max_input_chars", 1) < 100:
            raise HTTPException(status_code=422, detail="最大输入长度不能小于 100")
        row.safety_json = {**row.safety_json, **body.safety}
    if body.pricing:
        allowed_pricing = {"currency", "token_price_per_1k", "gpu_price_per_second", "cpu_price_per_second"}
        invalid = sorted(set(body.pricing) - allowed_pricing)
        if invalid:
            raise HTTPException(status_code=422, detail=f"未知计费字段: {', '.join(invalid)}")
        if any(float(value) < 0 for key, value in body.pricing.items() if key != "currency"):
            raise HTTPException(status_code=422, detail="计费价格不能为负数")
        row.pricing_json = {**row.pricing_json, **body.pricing}
    await audit(db, ctx, "tenant.policy.update", "tenant_policy", row.id, {"safety": body.safety, "pricing": body.pricing})
    await db.commit()
    return {"tenant_id": ctx.tenant_id, "safety": row.safety_json, "pricing": row.pricing_json}


@app.get("/api/admin/identity/status")
async def identity_status(ctx: AdminContext):
    settings = get_settings()
    return {
        "oidc": {
            "configured": bool(settings.oidc_issuer_url and settings.oidc_client_id and settings.oidc_client_secret),
            "issuer": settings.oidc_issuer_url,
            "auto_provision": settings.oidc_auto_provision,
            "sync_roles": settings.oidc_sync_roles,
            "groups_claim": settings.oidc_groups_claim,
            "roles_claim": settings.oidc_roles_claim,
            "default_role": settings.oidc_default_role,
            "mapped_groups": len(parse_role_mapping(settings.oidc_group_role_map)),
        },
        "ldap": {
            "configured": bool(settings.ldap_url and settings.ldap_user_search_base),
            "url": settings.ldap_url,
            "auto_provision": settings.ldap_auto_provision,
        },
    }


@app.post("/api/admin/alerts/test")
async def test_alert(body: dict[str, Any], ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    target = str(body.get("target") or "generic")
    result = await send_alert(str(body.get("title") or "企业 AI 中台告警测试"), str(body.get("text") or "告警通道连通性测试"), str(body.get("severity") or "info"), target)
    await audit(db, ctx, "alert.test", "alert_channel", target, result)
    await db.commit()
    return result


@app.post("/api/admin/security/rotate-secrets")
async def rotate_secrets(ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    rotated = {"models": 0, "plugins": 0, "mcp_servers": 0, "mfa": 0}
    for row in (await db.execute(select(ModelProvider).where(ModelProvider.tenant_id == ctx.tenant_id))).scalars().all():
        if row.api_key_encrypted:
            value = decrypt_secret(row.api_key_encrypted)
            if value:
                row.api_key_encrypted = encrypt_secret(value)
                rotated["models"] += 1
    for row in (await db.execute(select(Plugin).where(Plugin.tenant_id == ctx.tenant_id))).scalars().all():
        if row.secret_encrypted:
            value = decrypt_secret(row.secret_encrypted)
            if value:
                row.secret_encrypted = encrypt_secret(value)
                rotated["plugins"] += 1
    for row in (await db.execute(select(MCPServer).where(MCPServer.tenant_id == ctx.tenant_id))).scalars().all():
        if row.secret_encrypted:
            value = decrypt_secret(row.secret_encrypted)
            if value:
                row.secret_encrypted = encrypt_secret(value)
                rotated["mcp_servers"] += 1
    users = (await db.execute(select(User).join(Membership, Membership.user_id == User.id).where(Membership.tenant_id == ctx.tenant_id, User.mfa_secret_encrypted != ""))).scalars().all()
    for row in users:
        value = decrypt_secret(row.mfa_secret_encrypted)
        if value:
            row.mfa_secret_encrypted = encrypt_secret(value)
            rotated["mfa"] += 1
    await audit(db, ctx, "security.secrets.rotate", "tenant", ctx.tenant_id, rotated)
    await db.commit()
    return {"rotated": rotated, "backend": secret_backend_status()}


@app.post("/api/admin/scim/tokens")
async def create_scim_token(body: ScimTokenCreate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    raw = f"scim_{secrets.token_urlsafe(36)}"
    row = ScimToken(tenant_id=ctx.tenant_id, name=body.name, token_hash=hashlib.sha256(raw.encode()).hexdigest(), created_by=ctx.user_id)
    db.add(row)
    await db.flush()
    await audit(db, ctx, "scim.token.create", "scim_token", row.id, {"name": row.name})
    await db.commit()
    return {"id": row.id, "name": row.name, "token": raw, "warning": "令牌只在本次响应中显示"}


@app.get("/api/admin/scim/tokens")
async def list_scim_tokens(ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(ScimToken).where(ScimToken.tenant_id == ctx.tenant_id).order_by(ScimToken.created_at.desc()))).scalars().all()
    return [{"id": row.id, "name": row.name, "enabled": row.enabled, "last_used_at": row.last_used_at, "created_at": row.created_at} for row in rows]


@app.delete("/api/admin/scim/tokens/{token_id}")
async def revoke_scim_token(token_id: str, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(ScimToken).where(ScimToken.id == token_id, ScimToken.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="SCIM 令牌不存在")
    row.enabled = False
    await audit(db, ctx, "scim.token.revoke", "scim_token", row.id)
    await db.commit()
    return {"id": token_id, "enabled": False}


async def scim_context(tenant_slug: str, authorization: str | None, db: AsyncSession) -> tuple[Tenant, ScimToken]:
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(status_code=401, detail="SCIM 需要 Bearer 令牌")
    raw = authorization.split(" ", 1)[1].strip()
    tenant = await db.scalar(select(Tenant).where(Tenant.slug == tenant_slug, Tenant.is_active.is_(True)))
    token = await db.scalar(select(ScimToken).where(ScimToken.tenant_id == (tenant.id if tenant else ""), ScimToken.token_hash == hashlib.sha256(raw.encode()).hexdigest(), ScimToken.enabled.is_(True)))
    if not tenant or not token:
        raise HTTPException(status_code=401, detail="SCIM 令牌无效")
    token.last_used_at = utcnow()
    return tenant, token


def scim_user_payload(user: User, tenant_slug: str) -> dict[str, Any]:
    return {
        "schemas": ["urn:ietf:params:scim:schemas:core:2.0:User"],
        "id": user.id,
        "userName": user.email,
        "displayName": user.display_name,
        "active": user.is_active,
        "meta": {"resourceType": "User", "location": f"/scim/v2.0/{tenant_slug}/Users/{user.id}"},
    }


@app.get("/scim/v2.0/{tenant_slug}/Users")
async def scim_list_users(tenant_slug: str, authorization: str | None = Header(default=None), startIndex: int = 1, count: int = 100, db: AsyncSession = Depends(get_admin_db)):
    tenant, token = await scim_context(tenant_slug, authorization, db)
    users = (await db.execute(select(User).join(Membership, Membership.user_id == User.id).where(Membership.tenant_id == tenant.id).order_by(User.email).offset(max(0, startIndex - 1)).limit(min(1000, max(1, count))))).scalars().all()
    await db.commit()
    return {"schemas": ["urn:ietf:params:scim:api:messages:2.0:ListResponse"], "totalResults": len(users), "startIndex": startIndex, "itemsPerPage": len(users), "Resources": [scim_user_payload(user, tenant.slug) for user in users]}


@app.post("/scim/v2.0/{tenant_slug}/Users")
async def scim_create_user(tenant_slug: str, body: ScimUserCreate, authorization: str | None = Header(default=None), db: AsyncSession = Depends(get_admin_db)):
    tenant, token = await scim_context(tenant_slug, authorization, db)
    email = body.userName.lower().strip()
    user = await db.scalar(select(User).where(User.email == email))
    if user:
        raise HTTPException(status_code=409, detail="SCIM 用户已存在")
    user = User(email=email, display_name=body.displayName or email, password_hash=hash_password(secrets.token_urlsafe(32)), is_active=body.active)
    db.add(user)
    await db.flush()
    role = "viewer"
    for item in body.roles:
        value = str(item.get("value") or item.get("display") or "")
        if value in BUILTIN_ROLES:
            role = value
            break
    db.add(Membership(user_id=user.id, tenant_id=tenant.id, role=role))
    db.add(AuditEvent(tenant_id=tenant.id, actor_id=token.created_by, action="scim.user.create", resource_type="user", resource_id=user.id, detail_json={"email": email, "role": role}))
    await db.commit()
    return scim_user_payload(user, tenant.slug)


@app.get("/scim/v2.0/{tenant_slug}/Users/{user_id}")
async def scim_get_user(tenant_slug: str, user_id: str, authorization: str | None = Header(default=None), db: AsyncSession = Depends(get_admin_db)):
    tenant, token = await scim_context(tenant_slug, authorization, db)
    user = await db.scalar(select(User).join(Membership, Membership.user_id == User.id).where(User.id == user_id, Membership.tenant_id == tenant.id))
    if not user:
        raise HTTPException(status_code=404, detail="SCIM 用户不存在")
    await db.commit()
    return scim_user_payload(user, tenant.slug)


@app.patch("/scim/v2.0/{tenant_slug}/Users/{user_id}")
async def scim_patch_user(tenant_slug: str, user_id: str, body: ScimUserPatch, authorization: str | None = Header(default=None), db: AsyncSession = Depends(get_admin_db)):
    tenant, token = await scim_context(tenant_slug, authorization, db)
    user = await db.scalar(select(User).join(Membership, Membership.user_id == User.id).where(User.id == user_id, Membership.tenant_id == tenant.id))
    if not user:
        raise HTTPException(status_code=404, detail="SCIM 用户不存在")
    for operation in body.Operations:
        path = str(operation.get("path") or "").lower()
        value = operation.get("value")
        if path == "active":
            user.is_active = bool(value)
        elif path == "displayname":
            user.display_name = str(value or "")
        elif path == "username":
            user.email = str(value).lower().strip()
        elif isinstance(value, dict):
            if "active" in value:
                user.is_active = bool(value["active"])
            if "displayName" in value:
                user.display_name = str(value["displayName"])
    db.add(AuditEvent(tenant_id=tenant.id, actor_id=token.created_by, action="scim.user.update", resource_type="user", resource_id=user.id, detail_json={"operations": body.Operations}))
    await db.commit()
    return scim_user_payload(user, tenant.slug)


@app.delete("/scim/v2.0/{tenant_slug}/Users/{user_id}")
async def scim_delete_user(tenant_slug: str, user_id: str, authorization: str | None = Header(default=None), db: AsyncSession = Depends(get_admin_db)):
    tenant, token = await scim_context(tenant_slug, authorization, db)
    user = await db.scalar(select(User).join(Membership, Membership.user_id == User.id).where(User.id == user_id, Membership.tenant_id == tenant.id))
    if not user:
        raise HTTPException(status_code=404, detail="SCIM 用户不存在")
    user.is_active = False
    db.add(AuditEvent(tenant_id=tenant.id, actor_id=token.created_by, action="scim.user.disable", resource_type="user", resource_id=user.id, detail_json={}))
    await db.commit()
    return {"deleted": True, "id": user_id}


@app.get("/api/permissions")
async def list_permissions(ctx: PermissionAdminContext):
    return PERMISSION_CATALOG


@app.get("/api/roles")
async def list_roles(ctx: PermissionAdminContext, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(Role).where(Role.tenant_id == ctx.tenant_id).order_by(Role.name))).scalars().all()
    return [{"id": row.id, "name": row.name, "slug": row.slug, "description": row.description, "permissions": row.permissions_json, "is_system": row.is_system} for row in rows]


@app.post("/api/roles")
async def create_role(body: RoleCreate, ctx: PermissionAdminContext, db: AsyncSession = Depends(get_db)):
    allowed = {item["key"] for item in PERMISSION_CATALOG}
    invalid = sorted(set(body.permissions) - allowed)
    if invalid:
        raise HTTPException(status_code=422, detail=f"未知权限: {', '.join(invalid)}")
    row = Role(tenant_id=ctx.tenant_id, name=body.name, slug=body.slug, description=body.description, permissions_json=body.permissions, created_by=ctx.user_id)
    db.add(row)
    try:
        await db.flush()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(status_code=409, detail="角色标识已存在") from exc
    await audit(db, ctx, "role.create", "role", row.id, {"slug": row.slug, "permissions": row.permissions_json})
    await db.commit()
    return {"id": row.id, "name": row.name, "slug": row.slug, "permissions": row.permissions_json}


@app.patch("/api/roles/{role_id}")
async def update_role(role_id: str, body: RoleUpdate, ctx: PermissionAdminContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(Role).where(Role.id == role_id, Role.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="角色不存在")
    if row.is_system:
        raise HTTPException(status_code=422, detail="系统角色不可修改")
    if body.name is not None:
        row.name = body.name
    if body.description is not None:
        row.description = body.description
    if body.permissions is not None:
        allowed = {item["key"] for item in PERMISSION_CATALOG}
        invalid = sorted(set(body.permissions) - allowed)
        if invalid:
            raise HTTPException(status_code=422, detail=f"未知权限: {', '.join(invalid)}")
        row.permissions_json = body.permissions
    await audit(db, ctx, "role.update", "role", row.id, body.model_dump(exclude_none=True))
    await db.commit()
    return {"id": row.id, "name": row.name, "slug": row.slug, "description": row.description, "permissions": row.permissions_json}


@app.delete("/api/roles/{role_id}")
async def delete_role(role_id: str, ctx: PermissionAdminContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(Role).where(Role.id == role_id, Role.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="角色不存在")
    if row.is_system:
        raise HTTPException(status_code=422, detail="系统角色不可删除")
    assigned = await db.scalar(select(Membership).where(Membership.custom_role_id == row.id))
    if assigned:
        raise HTTPException(status_code=409, detail="角色仍被成员使用")
    await audit(db, ctx, "role.delete", "role", row.id)
    await db.delete(row)
    await db.commit()
    return {"id": role_id, "deleted": True}


@app.get("/api/tenants")
async def tenants(ctx: Context, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Membership, Tenant).join(Tenant, Tenant.id == Membership.tenant_id).where(Membership.user_id == ctx.user_id, Tenant.is_active.is_(True)))
    return [{"id": tenant.id, "name": tenant.name, "slug": tenant.slug, "role": membership.role} for membership, tenant in result.all()]


@app.get("/api/members")
async def list_members(ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(User, Membership).join(Membership, Membership.user_id == User.id).where(Membership.tenant_id == ctx.tenant_id).order_by(User.email))
    return [{"user_id": user.id, "email": user.email, "display_name": user.display_name, "role": membership.role, "custom_role_id": membership.custom_role_id, "is_active": user.is_active, "created_at": user.created_at} for user, membership in result.all()]


@app.post("/api/members")
async def create_member(body: MemberCreate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    user = await db.scalar(select(User).where(User.email == body.email.lower().strip()))
    if not user:
        user = User(email=body.email.lower().strip(), display_name=body.display_name, password_hash=hash_password(body.password))
        db.add(user)
        await db.flush()
    membership = await db.scalar(select(Membership).where(Membership.user_id == user.id, Membership.tenant_id == ctx.tenant_id))
    if membership:
        raise HTTPException(status_code=409, detail="该用户已经是当前租户成员")
    custom_role = None
    membership = Membership(user_id=user.id, tenant_id=ctx.tenant_id, role=body.role, custom_role_id=custom_role)
    db.add(membership)
    await db.flush()
    await audit(db, ctx, "member.create", "membership", membership.id, {"email": user.email, "role": body.role})
    await db.commit()
    return {"user_id": user.id, "email": user.email, "display_name": user.display_name, "role": membership.role, "is_active": user.is_active}


@app.patch("/api/members/{user_id}")
async def update_member(user_id: str, body: MemberRoleUpdate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    membership = await db.scalar(select(Membership).where(Membership.user_id == user_id, Membership.tenant_id == ctx.tenant_id))
    user = await db.scalar(select(User).where(User.id == user_id))
    if not membership or not user:
        raise HTTPException(status_code=404, detail="租户成员不存在")
    if user_id == ctx.user_id and body.is_active is False:
        raise HTTPException(status_code=422, detail="不能停用当前登录账号")
    membership.role = body.role
    if body.custom_role_id is None:
        membership.custom_role_id = None
    else:
        custom_role = await db.scalar(select(Role).where(Role.id == body.custom_role_id, Role.tenant_id == ctx.tenant_id))
        if not custom_role:
            raise HTTPException(status_code=404, detail="自定义角色不存在")
        membership.custom_role_id = custom_role.id
    if body.is_active is not None:
        user.is_active = body.is_active
    await audit(db, ctx, "member.update", "membership", membership.id, {"user_id": user_id, "role": body.role, "is_active": user.is_active})
    await db.commit()
    return {"user_id": user.id, "email": user.email, "display_name": user.display_name, "role": membership.role, "custom_role_id": membership.custom_role_id, "is_active": user.is_active}


@app.delete("/api/members/{user_id}")
async def remove_member(user_id: str, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    if user_id == ctx.user_id:
        raise HTTPException(status_code=422, detail="不能移除当前登录账号")
    membership = await db.scalar(select(Membership).where(Membership.user_id == user_id, Membership.tenant_id == ctx.tenant_id))
    if not membership:
        raise HTTPException(status_code=404, detail="租户成员不存在")
    await audit(db, ctx, "member.remove", "membership", membership.id, {"user_id": user_id})
    await db.delete(membership)
    await db.commit()
    return {"user_id": user_id, "tenant_id": ctx.tenant_id, "removed": True}


@app.get("/api/projects")
async def list_projects(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(Project).where(Project.tenant_id == ctx.tenant_id).order_by(Project.name))).scalars().all()
    return [{"id": row.id, "name": row.name, "slug": row.slug, "description": row.description} for row in rows]


@app.post("/api/projects")
async def create_project(body: ProjectCreate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    row = Project(tenant_id=ctx.tenant_id, name=body.name, slug=body.slug, description=body.description)
    db.add(row)
    try:
        await db.flush()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(status_code=409, detail="项目 slug 已存在") from exc
    await audit(db, ctx, "project.create", "project", row.id, {"slug": row.slug})
    await db.commit()
    return {"id": row.id, "name": row.name, "slug": row.slug, "description": row.description}


@app.get("/api/applications")
async def list_applications(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(Application).where(Application.tenant_id == ctx.tenant_id).order_by(Application.updated_at.desc()))).scalars().all()
    return [{"id": row.id, "project_id": row.project_id, "name": row.name, "slug": row.slug, "description": row.description, "app_type": row.app_type, "status": row.status, "runtime": row.runtime_json, "created_at": row.created_at, "updated_at": row.updated_at} for row in rows]


@app.post("/api/applications")
async def create_application(body: ApplicationCreate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    project = await db.scalar(select(Project).where(Project.id == body.project_id, Project.tenant_id == ctx.tenant_id))
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    row = Application(tenant_id=ctx.tenant_id, project_id=project.id, name=body.name, slug=body.slug, description=body.description, app_type=body.app_type, runtime_json=body.runtime, created_by=ctx.user_id)
    db.add(row)
    await db.flush()
    version = ApplicationVersion(tenant_id=ctx.tenant_id, application_id=row.id, version=1, source_json=body.source, manifest_json=body.runtime, created_by=ctx.user_id)
    db.add(version)
    await audit(db, ctx, "application.create", "application", row.id, {"app_type": row.app_type})
    await db.commit()
    return {"id": row.id, "name": row.name, "slug": row.slug, "app_type": row.app_type, "status": row.status, "version": 1}


@app.get("/api/applications/{application_id}")
async def get_application(application_id: str, ctx: Context, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(Application).where(Application.id == application_id, Application.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="应用不存在")
    versions = (await db.execute(select(ApplicationVersion).where(ApplicationVersion.application_id == row.id, ApplicationVersion.tenant_id == ctx.tenant_id).order_by(ApplicationVersion.version.desc()))).scalars().all()
    return {"id": row.id, "project_id": row.project_id, "name": row.name, "slug": row.slug, "description": row.description, "app_type": row.app_type, "status": row.status, "runtime": row.runtime_json, "versions": [{"id": item.id, "version": item.version, "source": item.source_json, "manifest": item.manifest_json, "status": item.status, "created_at": item.created_at} for item in versions]}


@app.post("/api/applications/{application_id}/versions")
async def create_application_version(application_id: str, body: ApplicationVersionCreate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    app_row = await db.scalar(select(Application).where(Application.id == application_id, Application.tenant_id == ctx.tenant_id))
    if not app_row:
        raise HTTPException(status_code=404, detail="应用不存在")
    latest = await db.scalar(select(func.max(ApplicationVersion.version)).where(ApplicationVersion.application_id == application_id, ApplicationVersion.tenant_id == ctx.tenant_id)) or 0
    version = ApplicationVersion(tenant_id=ctx.tenant_id, application_id=application_id, version=int(latest) + 1, source_json=body.source, manifest_json=body.manifest, created_by=ctx.user_id)
    db.add(version)
    app_row.status = "draft"
    await audit(db, ctx, "application.version.create", "application", application_id, {"version": version.version})
    await db.commit()
    return {"application_id": application_id, "version": version.version, "status": version.status}


@app.post("/api/applications/{application_id}/deploy")
async def deploy_application(application_id: str, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    app_row = await db.scalar(select(Application).where(Application.id == application_id, Application.tenant_id == ctx.tenant_id))
    if not app_row:
        raise HTTPException(status_code=404, detail="应用不存在")
    version = await db.scalar(select(ApplicationVersion).where(ApplicationVersion.application_id == application_id, ApplicationVersion.tenant_id == ctx.tenant_id).order_by(ApplicationVersion.version.desc()))
    if not version:
        raise HTTPException(status_code=422, detail="应用还没有版本")
    version.status = "deployed"
    app_row.status = "deployed"
    await audit(db, ctx, "application.deploy", "application", application_id, {"version": version.version})
    await db.commit()
    return {"application_id": application_id, "version": version.version, "status": "deployed"}


@app.patch("/api/applications/{application_id}")
async def update_application(application_id: str, body: ApplicationUpdate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(Application).where(Application.id == application_id, Application.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="应用不存在")
    for field in ("name", "description", "runtime"):
        value = getattr(body, field)
        if value is not None:
            setattr(row, "runtime_json" if field == "runtime" else field, value)
    row.status = "draft"
    await audit(db, ctx, "application.update", "application", row.id, body.model_dump(exclude_none=True))
    await db.commit()
    return {"id": row.id, "name": row.name, "description": row.description, "runtime": row.runtime_json, "status": row.status}


@app.post("/api/applications/{application_id}/rollback")
async def rollback_application(application_id: str, body: ApplicationRollback, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(Application).where(Application.id == application_id, Application.tenant_id == ctx.tenant_id))
    version = await db.scalar(select(ApplicationVersion).where(ApplicationVersion.application_id == application_id, ApplicationVersion.tenant_id == ctx.tenant_id, ApplicationVersion.version == body.version))
    if not row or not version:
        raise HTTPException(status_code=404, detail="应用或目标版本不存在")
    for item in (await db.execute(select(ApplicationVersion).where(ApplicationVersion.application_id == application_id))).scalars().all():
        item.status = "deployed" if item.id == version.id else "archived"
    row.status = "deployed"
    row.runtime_json = version.manifest_json or {}
    await audit(db, ctx, "application.rollback", "application", row.id, {"version": body.version})
    await db.commit()
    return {"application_id": row.id, "version": version.version, "status": "deployed"}


@app.delete("/api/applications/{application_id}")
async def delete_application(application_id: str, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(Application).where(Application.id == application_id, Application.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="应用不存在")
    await db.execute(delete(ApplicationVersion).where(ApplicationVersion.application_id == row.id, ApplicationVersion.tenant_id == ctx.tenant_id))
    await audit(db, ctx, "application.delete", "application", row.id)
    await db.delete(row)
    await db.commit()
    return {"id": application_id, "deleted": True}


@app.get("/api/applications/{application_id}/runtime")
async def application_runtime(application_id: str, ctx: Context, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(Application).where(Application.id == application_id, Application.tenant_id == ctx.tenant_id, Application.status == "deployed"))
    version = await db.scalar(select(ApplicationVersion).where(ApplicationVersion.application_id == application_id, ApplicationVersion.tenant_id == ctx.tenant_id, ApplicationVersion.status == "deployed").order_by(ApplicationVersion.version.desc()))
    if not row or not version:
        raise HTTPException(status_code=404, detail="应用没有已发布版本")
    return {"application": {"id": row.id, "name": row.name, "slug": row.slug, "type": row.app_type}, "version": version.version, "runtime": version.manifest_json, "source": version.source_json}


@app.get("/public/apps/{tenant_slug}/{app_slug}")
async def public_application_runtime(tenant_slug: str, app_slug: str, db: AsyncSession = Depends(get_admin_db)):
    tenant = await db.scalar(select(Tenant).where(Tenant.slug == tenant_slug, Tenant.is_active.is_(True)))
    row = await db.scalar(select(Application).where(Application.tenant_id == tenant.id if tenant else "", Application.slug == app_slug, Application.status == "deployed")) if tenant else None
    version = await db.scalar(select(ApplicationVersion).where(ApplicationVersion.application_id == row.id, ApplicationVersion.status == "deployed").order_by(ApplicationVersion.version.desc())) if row else None
    if not row or not version:
        raise HTTPException(status_code=404, detail="应用不存在或尚未发布")
    return {"name": row.name, "slug": row.slug, "type": row.app_type, "version": version.version, "runtime": version.manifest_json, "source": version.source_json}


@app.post("/public/apps/{tenant_slug}/{app_slug}/invoke")
async def invoke_public_application(
    tenant_slug: str,
    app_slug: str,
    body: PublicAppInvokeRequest,
    x_app_key: str | None = Header(default=None),
    db: AsyncSession = Depends(get_admin_db),
):
    tenant = await db.scalar(select(Tenant).where(Tenant.slug == tenant_slug, Tenant.is_active.is_(True)))
    row = await db.scalar(select(Application).where(Application.tenant_id == tenant.id if tenant else "", Application.slug == app_slug, Application.status == "deployed")) if tenant else None
    version = await db.scalar(select(ApplicationVersion).where(ApplicationVersion.application_id == row.id, ApplicationVersion.status == "deployed").order_by(ApplicationVersion.version.desc())) if row else None
    if not row or not version:
        raise HTTPException(status_code=404, detail="应用不存在或尚未发布")
    manifest = version.manifest_json or {}
    if manifest.get("public") is False:
        raise HTTPException(status_code=403, detail="应用未开放公开调用")
    expected_key = str(manifest.get("app_key") or (version.source_json or {}).get("public_key") or "")
    if expected_key and not x_app_key or expected_key and not hmac.compare_digest(expected_key, x_app_key or ""):
        raise HTTPException(status_code=401, detail="应用调用密钥无效")
    workflow_id = str(manifest.get("workflow_id") or "")
    if not workflow_id:
        if row.app_type == "full_code":
            raise HTTPException(status_code=501, detail="全代码应用需要配置沙箱运行时镜像")
        return {"status": "accepted", "application": row.slug, "version": version.version, "output": manifest.get("response", {})}
    workflow = await db.scalar(select(Workflow).where(Workflow.id == workflow_id, Workflow.tenant_id == tenant.id, Workflow.status == "published"))
    if not workflow or not workflow.published_version:
        raise HTTPException(status_code=422, detail="应用绑定的工作流尚未发布")
    run = WorkflowRun(tenant_id=tenant.id, project_id=workflow.project_id, workflow_id=workflow.id, workflow_version=workflow.published_version, input_json=body.input, trace_id="", created_by=row.created_by)
    db.add(run)
    await db.flush()
    await db.commit()
    await enqueue_run(run.id)
    return {"status": "queued", "run_id": run.id, "application": row.slug, "version": version.version, "workflow_id": workflow.id}


@app.get("/api/compute/quota")
async def get_compute_quota(ctx: Context, db: AsyncSession = Depends(get_db)):
    quota = await db.scalar(select(ComputeQuota).where(ComputeQuota.tenant_id == ctx.tenant_id))
    if not quota:
        quota = ComputeQuota(tenant_id=ctx.tenant_id)
        db.add(quota)
        await db.commit()
    return {"tenant_id": ctx.tenant_id, "cpu_millicores": quota.cpu_millicores, "memory_mb": quota.memory_mb, "gpu_count": quota.gpu_count, "concurrent_runs": quota.concurrent_runs, "monthly_tokens": quota.monthly_tokens, "enabled": quota.enabled}


@app.put("/api/compute/quota")
async def update_compute_quota(body: QuotaUpdate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    quota = await db.scalar(select(ComputeQuota).where(ComputeQuota.tenant_id == ctx.tenant_id))
    if not quota:
        quota = ComputeQuota(tenant_id=ctx.tenant_id)
        db.add(quota)
    for field in ("cpu_millicores", "memory_mb", "gpu_count", "concurrent_runs", "monthly_tokens", "enabled"):
        setattr(quota, field, getattr(body, field))
    await audit(db, ctx, "compute.quota.update", "compute_quota", quota.id, body.model_dump())
    await db.commit()
    return {"tenant_id": ctx.tenant_id, **body.model_dump()}


@app.get("/api/compute/usage")
async def get_compute_usage(ctx: Context, db: AsyncSession = Depends(get_db)):
    row = (await db.execute(select(func.coalesce(func.sum(ComputeUsage.cpu_milliseconds), 0), func.coalesce(func.sum(ComputeUsage.memory_mb_seconds), 0), func.coalesce(func.sum(ComputeUsage.gpu_seconds), 0), func.coalesce(func.sum(ComputeUsage.tokens), 0)).where(ComputeUsage.tenant_id == ctx.tenant_id))).one()
    return {"tenant_id": ctx.tenant_id, "cpu_milliseconds": int(row[0]), "memory_mb_seconds": int(row[1]), "gpu_seconds": int(row[2]), "tokens": int(row[3])}


async def reconcile_gpu_leases(db: AsyncSession, tenant_id: str | None = None) -> int:
    now_value = utcnow()
    stmt = select(GpuLease).where(GpuLease.status.in_(["allocated", "queued"]), GpuLease.expires_at <= now_value)
    if tenant_id:
        stmt = stmt.where(GpuLease.tenant_id == tenant_id)
    expired = (await db.execute(stmt)).scalars().all()
    for lease in expired:
        lease.status = "expired"
        lease.released_at = now_value
    return len(expired)


@app.get("/api/compute/gpu/leases")
async def list_gpu_leases(ctx: Context, db: AsyncSession = Depends(get_db)):
    await reconcile_gpu_leases(db, ctx.tenant_id)
    rows = (await db.execute(select(GpuLease).where(GpuLease.tenant_id == ctx.tenant_id).order_by(desc(GpuLease.created_at)).limit(100))).scalars().all()
    await db.commit()
    return [{"id": row.id, "requested_count": row.requested_count, "status": row.status, "node_name": row.node_name, "run_id": row.run_id, "expires_at": row.expires_at, "created_at": row.created_at} for row in rows]


@app.post("/api/compute/gpu/leases")
async def create_gpu_lease(body: GpuLeaseRequest, ctx: OperatorContext, db: AsyncSession = Depends(get_db)):
    settings = get_settings()
    if settings.gpu_total_count < body.requested_count:
        raise HTTPException(status_code=503, detail="平台 GPU 总容量未配置或不足")
    await reconcile_gpu_leases(db, ctx.tenant_id)
    active = await db.scalar(select(func.coalesce(func.sum(GpuLease.requested_count), 0)).where(GpuLease.status == "allocated", GpuLease.expires_at > utcnow())) or 0
    status = "allocated" if int(active) + body.requested_count <= settings.gpu_total_count else "queued"
    row = GpuLease(tenant_id=ctx.tenant_id, run_id=body.run_id, requested_count=body.requested_count, status=status, node_name="scheduler-pending" if status == "queued" else "unassigned", expires_at=utcnow() + timedelta(seconds=min(body.ttl_seconds, settings.gpu_lease_ttl_seconds)), created_by=ctx.user_id)
    db.add(row)
    await audit(db, ctx, "compute.gpu.lease.create", "gpu_lease", row.id, {"count": body.requested_count, "status": status})
    await db.commit()
    return {"id": row.id, "status": row.status, "requested_count": row.requested_count, "expires_at": row.expires_at, "capacity": settings.gpu_total_count, "allocated": int(active) + (body.requested_count if status == "allocated" else 0)}


@app.delete("/api/compute/gpu/leases/{lease_id}")
async def release_gpu_lease(lease_id: str, ctx: OperatorContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(GpuLease).where(GpuLease.id == lease_id, GpuLease.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="GPU 租约不存在")
    row.status = "released"
    row.released_at = utcnow()
    await audit(db, ctx, "compute.gpu.lease.release", "gpu_lease", row.id)
    await db.commit()
    return {"id": row.id, "status": row.status}


@app.get("/api/finops/summary")
async def finops_summary(ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    usage = await get_compute_usage(ctx, db)
    policy = await tenant_policy(db, ctx.tenant_id)
    pricing = policy.pricing_json or {}
    token_cost = usage["tokens"] / 1000 * float(pricing.get("token_price_per_1k", get_settings().finops_token_price_per_1k))
    gpu_cost = usage["gpu_seconds"] * float(pricing.get("gpu_price_per_second", get_settings().finops_gpu_price_per_second))
    cpu_cost = usage["cpu_milliseconds"] / 1000 * float(pricing.get("cpu_price_per_second", get_settings().finops_cpu_price_per_second))
    await db.commit()
    return {
        "tenant_id": ctx.tenant_id,
        "currency": pricing.get("currency", get_settings().finops_currency),
        "usage": usage,
        "pricing": pricing,
        "estimated_cost": {"tokens": round(token_cost, 6), "gpu": round(gpu_cost, 6), "cpu": round(cpu_cost, 6), "total": round(token_cost + gpu_cost + cpu_cost, 6)},
        "budget": {"monthly_budget": float(pricing.get("monthly_budget", 0)), "exceeded": bool(pricing.get("monthly_budget", 0) and token_cost + gpu_cost + cpu_cost >= float(pricing.get("monthly_budget", 0)))},
    }


@app.put("/api/finops/budget")
async def update_finops_budget(body: BudgetUpdate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    policy = await tenant_policy(db, ctx.tenant_id)
    pricing = dict(policy.pricing_json or {})
    pricing["monthly_budget"] = body.monthly_budget
    policy.pricing_json = pricing
    await audit(db, ctx, "finops.budget.update", "tenant_policy", policy.id, {"monthly_budget": body.monthly_budget})
    await db.commit()
    return {"monthly_budget": body.monthly_budget, "currency": pricing.get("currency", get_settings().finops_currency)}


@app.get("/api/models")
async def list_models(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(ModelProvider).where(ModelProvider.tenant_id == ctx.tenant_id).order_by(ModelProvider.name))).scalars().all()
    return [{"id": row.id, "name": row.name, "provider_type": row.provider_type, "base_url": row.base_url, "default_model": row.default_model, "enabled": row.enabled, "has_api_key": bool(row.api_key_encrypted), "created_at": row.created_at} for row in rows]


@app.post("/api/models")
async def create_model(body: ProviderCreate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = ModelProvider(tenant_id=ctx.tenant_id, name=body.name, provider_type=body.provider_type, base_url=body.base_url.rstrip("/"), default_model=body.default_model, api_key_encrypted=encrypt_secret(body.api_key), config_json=body.config)
    db.add(row)
    try:
        await db.flush()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(status_code=409, detail="模型提供商名称已存在") from exc
    await audit(db, ctx, "model_provider.create", "model_provider", row.id, {"name": row.name})
    await db.commit()
    return {"id": row.id, "name": row.name, "provider_type": row.provider_type, "base_url": row.base_url, "default_model": row.default_model, "enabled": row.enabled, "has_api_key": bool(row.api_key_encrypted)}


@app.post("/api/models/{provider_id}/test")
async def test_model(provider_id: str, body: ModelTestRequest, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    provider = await db.scalar(select(ModelProvider).where(ModelProvider.id == provider_id, ModelProvider.tenant_id == ctx.tenant_id))
    if not provider:
        raise HTTPException(status_code=404, detail="模型提供商不存在")
    headers = {"Content-Type": "application/json"}
    api_key = decrypt_secret(provider.api_key_encrypted)
    if api_key:
        headers["Authorization"] = f"Bearer {api_key}"
    started = time.perf_counter()
    ok = False
    status_code = 0
    detail = ""
    try:
        if body.kind == "embedding":
            vectors = await embed(provider, [body.prompt])
            ok = bool(vectors and vectors[0])
            status_code = 200 if ok else 502
            detail = f"Embedding 接口可用，维度 {len(vectors[0])}" if ok else "Embedding 接口没有返回向量"
        else:
            endpoint = provider.base_url.rstrip("/")
            if body.kind == "chat":
                if not endpoint.endswith("/chat/completions"):
                    endpoint = f"{endpoint}/chat/completions"
                async with httpx.AsyncClient(timeout=get_settings().model_timeout_seconds) as client:
                    response = await client.post(endpoint, headers=headers, json={"model": provider.default_model, "messages": [{"role": "user", "content": body.prompt}], "max_tokens": 8})
            else:
                if not endpoint.endswith("/models"):
                    endpoint = f"{endpoint}/models"
                async with httpx.AsyncClient(timeout=get_settings().model_timeout_seconds) as client:
                    response = await client.get(endpoint, headers=headers)
            status_code = response.status_code
            ok = response.is_success
            detail = ("聊天接口可用" if body.kind == "chat" else "模型接口可用") if ok else response.text[:500]
    except Exception as exc:
        detail = str(exc)[:500]
    latency_ms = int((time.perf_counter() - started) * 1000)
    check = ModelHealthCheck(tenant_id=ctx.tenant_id, provider_id=provider.id, ok=ok, latency_ms=latency_ms, status_code=status_code, detail=detail)
    db.add(check)
    await audit(db, ctx, "model_provider.health_check", "model_provider", provider.id, {"ok": ok, "latency_ms": latency_ms, "status_code": status_code})
    await db.commit()
    return {"provider_id": provider.id, "ok": ok, "latency_ms": latency_ms, "status_code": status_code, "detail": detail}


@app.get("/api/models/{provider_id}/health")
async def model_health(provider_id: str, ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(ModelHealthCheck).where(ModelHealthCheck.provider_id == provider_id, ModelHealthCheck.tenant_id == ctx.tenant_id).order_by(desc(ModelHealthCheck.created_at)).limit(20))).scalars().all()
    return [{"ok": row.ok, "latency_ms": row.latency_ms, "status_code": row.status_code, "detail": row.detail, "created_at": row.created_at} for row in rows]


@app.get("/api/models/route")
async def route_model(ctx: Context, db: AsyncSession = Depends(get_db)):
    providers = (await db.execute(select(ModelProvider).where(ModelProvider.tenant_id == ctx.tenant_id, ModelProvider.enabled.is_(True)).order_by(ModelProvider.created_at))).scalars().all()
    candidates = []
    for provider in providers:
        health = await db.scalar(select(ModelHealthCheck).where(ModelHealthCheck.provider_id == provider.id, ModelHealthCheck.tenant_id == ctx.tenant_id).order_by(desc(ModelHealthCheck.created_at)))
        candidates.append({"id": provider.id, "name": provider.name, "model": provider.default_model, "healthy": health.ok if health else None, "latency_ms": health.latency_ms if health else None, "priority": (provider.config_json or {}).get("priority", 100)})
    candidates.sort(key=lambda item: (not item["healthy"] if item["healthy"] is not None else False, item["priority"], item["latency_ms"] or 10**9))
    return {"selected": candidates[0] if candidates else None, "candidates": candidates}


@app.get("/api/agents")
async def list_agents(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(Agent).where(Agent.tenant_id == ctx.tenant_id).order_by(Agent.name))).scalars().all()
    result = []
    for row in rows:
        config = await db.scalar(select(AgentConfig).where(AgentConfig.agent_id == row.id))
        result.append({"id": row.id, "project_id": row.project_id, "name": row.name, "slug": row.slug, "description": row.description, "provider_id": row.provider_id, "knowledge_base_id": row.knowledge_base_id, "enabled": row.enabled, "version": row.version, "agent_type": config.agent_type if config else "chat", "publish_status": config.publish_status if config else "draft", "workspace_id": config.workspace_id if config else None})
    return result


@app.post("/api/agents")
async def create_agent(body: AgentCreate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    project = await db.scalar(select(Project).where(Project.id == body.project_id, Project.tenant_id == ctx.tenant_id))
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    if body.provider_id and not await db.scalar(select(ModelProvider).where(ModelProvider.id == body.provider_id, ModelProvider.tenant_id == ctx.tenant_id)):
        raise HTTPException(status_code=404, detail="模型提供商不存在")
    if body.knowledge_base_id and not await db.scalar(select(KnowledgeBase).where(KnowledgeBase.id == body.knowledge_base_id, KnowledgeBase.tenant_id == ctx.tenant_id)):
        raise HTTPException(status_code=404, detail="知识库不存在")
    workspace_id = None
    if body.workspace_id:
        if not await db.scalar(select(Workspace).where(Workspace.id == body.workspace_id, Workspace.tenant_id == ctx.tenant_id)):
            raise HTTPException(status_code=404, detail="工作空间不存在")
        workspace_id = body.workspace_id
    row = Agent(tenant_id=ctx.tenant_id, project_id=project.id, name=body.name, slug=body.slug, description=body.description, system_prompt=body.system_prompt, provider_id=body.provider_id, knowledge_base_id=body.knowledge_base_id, created_by=ctx.user_id)
    db.add(row)
    try:
        await db.flush()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(status_code=409, detail="Agent slug 已存在") from exc
    db.add(AgentConfig(tenant_id=ctx.tenant_id, agent_id=row.id, workspace_id=workspace_id, agent_type="chat", prompt_template=body.system_prompt, experience_json={"welcome": "你好，我可以帮助你处理企业问题。", "suggested_questions": []}, safety_json={"enabled": False}, capabilities_json={"tools": [], "skills": [], "workflows": [], "agents": [], "knowledge": [], "databases": []}))
    await audit(db, ctx, "agent.create", "agent", row.id, {"name": row.name})
    await db.commit()
    return {"id": row.id, "name": row.name, "slug": row.slug, "version": row.version, "enabled": row.enabled}


@app.post("/api/agents/import")
async def import_agent(body: dict[str, Any], ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    project_id = str(body.get("project_id") or "")
    payload = body.get("agent") if isinstance(body.get("agent"), dict) else body
    project = await db.scalar(select(Project).where(Project.id == project_id, Project.tenant_id == ctx.tenant_id))
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    name = str(payload.get("name") or "导入 Agent")[:160]
    slug = str(payload.get("slug") or f"imported-{secrets.token_hex(4)}")[:100]
    slug = "".join(char if char.isalnum() or char == "-" else "-" for char in slug.lower()).strip("-") or f"imported-{secrets.token_hex(4)}"
    base = slug
    suffix = 1
    while await db.scalar(select(Agent).where(Agent.tenant_id == ctx.tenant_id, Agent.project_id == project.id, Agent.slug == slug)):
        suffix += 1
        slug = f"{base}-{suffix}"
    row = Agent(tenant_id=ctx.tenant_id, project_id=project.id, name=name, slug=slug, description=str(payload.get("description") or ""), system_prompt=str(payload.get("system_prompt") or "你是一名企业智能助手。"), provider_id=payload.get("provider_id"), knowledge_base_id=payload.get("knowledge_base_id"), created_by=ctx.user_id)
    db.add(row)
    await db.flush()
    config = body.get("config") if isinstance(body.get("config"), dict) else {}
    db.add(AgentConfig(tenant_id=ctx.tenant_id, agent_id=row.id, agent_type=str(config.get("agent_type") or "chat"), prompt_mode=str(config.get("prompt_mode") or "normal"), prompt_template=str(config.get("prompt_template") or row.system_prompt), variables_json=config.get("variables") or [], memory_json=config.get("memory") or {}, capabilities_json=config.get("capabilities") or {}, experience_json=config.get("experience") or {}, safety_json=config.get("safety") or {}))
    await audit(db, ctx, "agent.import", "agent", row.id, {"source_slug": payload.get("slug")})
    await db.commit()
    return {"id": row.id, "name": row.name, "slug": row.slug, "version": row.version, "imported": True}


@app.get("/api/agents/{agent_id}/export")
async def export_agent(agent_id: str, ctx: Context, db: AsyncSession = Depends(get_db)):
    agent = await db.scalar(select(Agent).where(Agent.id == agent_id, Agent.tenant_id == ctx.tenant_id))
    if not agent:
        raise HTTPException(status_code=404, detail="Agent 不存在")
    config = await db.scalar(select(AgentConfig).where(AgentConfig.agent_id == agent.id))
    versions = (await db.execute(select(AgentVersion).where(AgentVersion.agent_id == agent.id).order_by(AgentVersion.version))).scalars().all()
    return {"format": "enterprise-ai-agent", "version": 1, "agent": {"name": agent.name, "slug": agent.slug, "description": agent.description, "system_prompt": agent.system_prompt, "provider_id": agent.provider_id, "knowledge_base_id": agent.knowledge_base_id}, "config": {"agent_type": config.agent_type if config else "chat", "prompt_mode": config.prompt_mode if config else "normal", "prompt_template": config.prompt_template if config else agent.system_prompt, "variables": config.variables_json if config else [], "memory": config.memory_json if config else {}, "capabilities": config.capabilities_json if config else {}, "experience": config.experience_json if config else {}, "safety": config.safety_json if config else {}}, "versions": [{"version": item.version, "changelog": item.changelog, "config": item.config_json, "status": item.status} for item in versions]}


@app.post("/api/agents/{agent_id}/clone")
async def clone_agent(agent_id: str, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    source = await db.scalar(select(Agent).where(Agent.id == agent_id, Agent.tenant_id == ctx.tenant_id))
    if not source:
        raise HTTPException(status_code=404, detail="Agent 不存在")
    slug = f"{source.slug}-copy-{secrets.token_hex(3)}"
    clone = Agent(tenant_id=ctx.tenant_id, project_id=source.project_id, name=f"{source.name} 副本", slug=slug, description=source.description, system_prompt=source.system_prompt, provider_id=source.provider_id, knowledge_base_id=source.knowledge_base_id, created_by=ctx.user_id)
    db.add(clone)
    await db.flush()
    config = await db.scalar(select(AgentConfig).where(AgentConfig.agent_id == source.id))
    if config:
        db.add(AgentConfig(tenant_id=ctx.tenant_id, agent_id=clone.id, workspace_id=config.workspace_id, agent_type=config.agent_type, logo_url=config.logo_url, prompt_mode=config.prompt_mode, prompt_template=config.prompt_template, variables_json=config.variables_json, memory_json=config.memory_json, capabilities_json=config.capabilities_json, experience_json=config.experience_json, safety_json=config.safety_json))
    await audit(db, ctx, "agent.clone", "agent", clone.id, {"source_id": source.id})
    await db.commit()
    return {"id": clone.id, "name": clone.name, "slug": clone.slug, "version": clone.version}


@app.delete("/api/agents/{agent_id}")
async def delete_agent(agent_id: str, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    agent = await db.scalar(select(Agent).where(Agent.id == agent_id, Agent.tenant_id == ctx.tenant_id))
    if not agent:
        raise HTTPException(status_code=404, detail="Agent 不存在")
    await db.execute(delete(AgentEvent).where(AgentEvent.agent_id == agent.id, AgentEvent.tenant_id == ctx.tenant_id))
    await db.execute(delete(AgentVersion).where(AgentVersion.agent_id == agent.id, AgentVersion.tenant_id == ctx.tenant_id))
    await db.execute(delete(AgentConfig).where(AgentConfig.agent_id == agent.id, AgentConfig.tenant_id == ctx.tenant_id))
    await audit(db, ctx, "agent.delete", "agent", agent.id, {"name": agent.name})
    await db.delete(agent)
    await db.commit()
    return {"id": agent_id, "deleted": True}


@app.post("/api/agents/{agent_id}/chat")
async def chat_agent(agent_id: str, body: AgentChatRequest, ctx: OperatorContext, db: AsyncSession = Depends(get_db)):
    agent = await db.scalar(select(Agent).where(Agent.id == agent_id, Agent.tenant_id == ctx.tenant_id, Agent.enabled.is_(True)))
    if not agent:
        raise HTTPException(status_code=404, detail="Agent 不存在或已禁用")
    provider_stmt = select(ModelProvider).where(ModelProvider.tenant_id == ctx.tenant_id, ModelProvider.enabled.is_(True))
    if agent.provider_id:
        provider_stmt = provider_stmt.where(ModelProvider.id == agent.provider_id)
    provider = (await db.execute(provider_stmt)).scalars().first()
    if not provider:
        raise HTTPException(status_code=422, detail="Agent 没有绑定可用模型提供商")
    quota = await db.scalar(select(ComputeQuota).where(ComputeQuota.tenant_id == ctx.tenant_id))
    month_start = datetime(utcnow().year, utcnow().month, 1, tzinfo=timezone.utc)
    used_tokens = await db.scalar(select(func.coalesce(func.sum(ComputeUsage.tokens), 0)).where(ComputeUsage.tenant_id == ctx.tenant_id, ComputeUsage.created_at >= month_start)) or 0
    if quota and quota.enabled and int(used_tokens) >= quota.monthly_tokens:
        raise HTTPException(status_code=429, detail="已达到租户月度 Token 配额")
    context_text = ""
    if agent.knowledge_base_id:
        rows = await search_chunks(db, ctx.tenant_id, agent.knowledge_base_id, body.message, 5)
        context_text = "\n\n".join(row["content"] for row in rows)
    policy = await tenant_policy(db, ctx.tenant_id)
    safety = inspect_text(body.message, policy.safety_json)
    if safety["blocked"]:
        await audit(db, ctx, "safety.block", "agent", agent.id, {"reasons": safety["reasons"]})
        await db.commit()
        raise HTTPException(status_code=422, detail={"message": "输入未通过安全策略", "reasons": safety["reasons"]})
    user_message = safety["text"] if policy.safety_json.get("mode") == "redact" else body.message
    system = agent.system_prompt + (f"\n\nEnterprise knowledge context:\n{context_text}" if context_text else "")
    messages = [{"role": "system", "content": system}, *body.conversation[-20:], {"role": "user", "content": user_message}]
    started = time.perf_counter()
    result = await call_chat(db, provider, messages, tenant_id=ctx.tenant_id)
    usage = result.get("usage") or {}
    total_tokens = int(usage.get("total_tokens", 0))
    db.add(ComputeUsage(tenant_id=ctx.tenant_id, tokens=total_tokens))
    event = AgentEvent(tenant_id=ctx.tenant_id, agent_id=agent.id, user_id=ctx.user_id, event_type="message", tokens=total_tokens, latency_ms=int((time.perf_counter() - started) * 1000), detail_json={"provider_id": provider.id, "model": provider.default_model})
    db.add(event)
    await audit(db, ctx, "agent.chat", "agent", agent.id, {"provider": provider.name, "usage": result.get("usage", {})})
    await db.commit()
    await db.flush()
    return {"agent_id": agent.id, "event_id": event.id, "text": result["text"], "usage": result.get("usage", {})}


@app.post("/api/v1/agents/{agent_id}/chat")
async def public_chat_agent(agent_id: str, body: AgentChatRequest, ctx: OperatorContext, db: AsyncSession = Depends(get_db)):
    """API Key/JWT 对外 Agent 调用入口，复用内部租户和安全策略。"""
    require_api_scope(ctx, "agent:invoke")
    return await chat_agent(agent_id, body, ctx, db)


@app.post("/api/agents/{agent_id}/events/{event_id}/feedback")
async def agent_feedback(agent_id: str, event_id: str, body: AgentFeedbackCreate, ctx: OperatorContext, db: AsyncSession = Depends(get_db)):
    event = await db.scalar(select(AgentEvent).where(AgentEvent.id == event_id, AgentEvent.agent_id == agent_id, AgentEvent.tenant_id == ctx.tenant_id))
    if not event:
        raise HTTPException(status_code=404, detail="对话事件不存在")
    event.feedback = body.feedback
    event.detail_json = {**event.detail_json, "feedback_detail": body.detail}
    await audit(db, ctx, "agent.feedback", "agent_event", event.id, {"feedback": body.feedback})
    await db.commit()
    return {"event_id": event.id, "feedback": event.feedback}


@app.post("/api/safety/check")
async def safety_check(body: dict[str, Any], ctx: OperatorContext, db: AsyncSession = Depends(get_db)):
    policy = await tenant_policy(db, ctx.tenant_id)
    result = inspect_text(str(body.get("text") or ""), policy.safety_json)
    await db.commit()
    return result


@app.post("/api/evaluations/run")
async def run_evaluation(body: EvaluationRequest, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    agent = await db.scalar(select(Agent).where(Agent.id == body.agent_id, Agent.tenant_id == ctx.tenant_id, Agent.enabled.is_(True)))
    if not agent:
        raise HTTPException(status_code=404, detail="智能体不存在或已停用")
    provider_stmt = select(ModelProvider).where(ModelProvider.tenant_id == ctx.tenant_id, ModelProvider.enabled.is_(True))
    if agent.provider_id:
        provider_stmt = provider_stmt.where(ModelProvider.id == agent.provider_id)
    provider = (await db.execute(provider_stmt)).scalars().first()
    if not provider:
        raise HTTPException(status_code=422, detail="智能体没有绑定可用模型提供商")
    policy = await tenant_policy(db, ctx.tenant_id)
    results: list[dict[str, Any]] = []
    total_tokens = 0
    for case in body.cases:
        safety = inspect_text(case.message, policy.safety_json)
        if safety["blocked"]:
            results.append({"message": case.message, "passed": False, "score": 0, "error": "安全策略拦截", "reasons": safety["reasons"]})
            continue
        response = await call_chat(db, provider, [{"role": "system", "content": agent.system_prompt}, {"role": "user", "content": safety["text"]}], tenant_id=ctx.tenant_id)
        answer = response["text"]
        usage = response.get("usage") or {}
        total_tokens += int(usage.get("total_tokens", 0))
        missing = [expected for expected in case.expected_contains if expected.lower() not in answer.lower()]
        results.append({"message": case.message, "answer": answer, "passed": not missing, "score": 1 if not missing else 0, "missing": missing, "usage": usage})
    db.add(ComputeUsage(tenant_id=ctx.tenant_id, tokens=total_tokens))
    passed = sum(1 for item in results if item["passed"])
    await audit(db, ctx, "evaluation.run", "agent", agent.id, {"cases": len(results), "passed": passed, "tokens": total_tokens})
    await db.commit()
    return {"agent_id": agent.id, "total": len(results), "passed": passed, "score": round(passed / len(results), 4) if results else 0, "results": results}


@app.post("/api/evaluations/jobs")
async def create_evaluation_job(body: EvaluationJobCreate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    agent = await db.scalar(select(Agent).where(Agent.id == body.agent_id, Agent.tenant_id == ctx.tenant_id))
    if not agent:
        raise HTTPException(status_code=404, detail="智能体不存在")
    row = EvaluationJob(tenant_id=ctx.tenant_id, agent_id=body.agent_id, request_json=body.model_dump(), created_by=ctx.user_id)
    db.add(row)
    await db.flush()
    await audit(db, ctx, "evaluation.job.create", "evaluation_job", row.id, {"agent_id": body.agent_id, "cases": len(body.cases)})
    await db.commit()
    redis = Redis.from_url(get_settings().redis_url, decode_responses=True)
    await redis.rpush(get_settings().evaluation_queue, row.id)
    await redis.aclose()
    return {"id": row.id, "status": row.status, "created_at": row.created_at}


@app.get("/api/evaluations/jobs")
async def list_evaluation_jobs(ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(EvaluationJob).where(EvaluationJob.tenant_id == ctx.tenant_id).order_by(desc(EvaluationJob.created_at)).limit(100))).scalars().all()
    return [{"id": row.id, "agent_id": row.agent_id, "status": row.status, "results": row.results_json, "error": row.error, "created_at": row.created_at, "started_at": row.started_at, "finished_at": row.finished_at} for row in rows]


@app.get("/api/evaluations/jobs/{job_id}")
async def get_evaluation_job(job_id: str, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(EvaluationJob).where(EvaluationJob.id == job_id, EvaluationJob.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="评测任务不存在")
    return {"id": row.id, "agent_id": row.agent_id, "status": row.status, "request": row.request_json, "results": row.results_json, "error": row.error, "created_at": row.created_at, "started_at": row.started_at, "finished_at": row.finished_at}


@app.post("/api/runtime/sandbox/runs")
async def create_sandbox_run(body: SandboxRunCreate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    try:
        manifest = build_job_manifest("pending", body.runtime, body.source, body.timeout_seconds, body.cpu_millicores, body.memory_mb)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    row = SandboxRun(tenant_id=ctx.tenant_id, application_id=body.application_id, runtime=body.runtime, source_json=body.model_dump(), manifest_json=manifest, created_by=ctx.user_id)
    db.add(row)
    await db.flush()
    row.manifest_json = build_job_manifest(row.id, body.runtime, body.source, body.timeout_seconds, body.cpu_millicores, body.memory_mb)
    await audit(db, ctx, "sandbox.run.create", "sandbox_run", row.id, {"runtime": body.runtime})
    await db.commit()
    redis = Redis.from_url(get_settings().redis_url, decode_responses=True)
    await redis.rpush(get_settings().sandbox_queue, row.id)
    await redis.aclose()
    return {"id": row.id, "status": row.status, "runtime": row.runtime, "manifest": row.manifest_json, "created_at": row.created_at}


@app.get("/api/runtime/sandbox/runs")
async def list_sandbox_runs(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(SandboxRun).where(SandboxRun.tenant_id == ctx.tenant_id).order_by(desc(SandboxRun.created_at)).limit(100))).scalars().all()
    return [{"id": row.id, "status": row.status, "runtime": row.runtime, "output": row.output_json, "error": row.error, "created_at": row.created_at, "finished_at": row.finished_at} for row in rows]


@app.get("/api/runtime/sandbox/runs/{run_id}")
async def get_sandbox_run(run_id: str, ctx: Context, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(SandboxRun).where(SandboxRun.id == run_id, SandboxRun.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="沙箱任务不存在")
    return {"id": row.id, "status": row.status, "runtime": row.runtime, "source": row.source_json, "manifest": row.manifest_json, "output": row.output_json, "error": row.error, "created_at": row.created_at, "started_at": row.started_at, "finished_at": row.finished_at}


@app.post("/api/runtime/sandbox/runs/{run_id}/cancel")
async def cancel_sandbox_run(run_id: str, ctx: OperatorContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(SandboxRun).where(SandboxRun.id == run_id, SandboxRun.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="沙箱任务不存在")
    if row.status in {"queued", "running"}:
        row.status = "cancelled"
        row.finished_at = utcnow()
    await audit(db, ctx, "sandbox.run.cancel", "sandbox_run", row.id)
    await db.commit()
    return {"id": row.id, "status": row.status}


@app.get("/api/workflows")
async def list_workflows(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(Workflow).where(Workflow.tenant_id == ctx.tenant_id).order_by(desc(Workflow.updated_at)))).scalars().all()
    return [{"id": row.id, "project_id": row.project_id, "name": row.name, "slug": row.slug, "description": row.description, "status": row.status, "latest_version": row.latest_version, "published_version": row.published_version, "updated_at": row.updated_at} for row in rows]


@app.post("/api/workflows")
async def create_workflow(body: WorkflowCreate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    project = await db.scalar(select(Project).where(Project.id == body.project_id, Project.tenant_id == ctx.tenant_id))
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    try:
        ordered_nodes(body.graph)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    workflow = Workflow(tenant_id=ctx.tenant_id, project_id=project.id, name=body.name, slug=body.slug, description=body.description, latest_version=1, created_by=ctx.user_id)
    db.add(workflow)
    await db.flush()
    db.add(WorkflowVersion(workflow_id=workflow.id, tenant_id=ctx.tenant_id, version=1, graph_json=body.graph, input_schema=body.input_schema, changelog=body.changelog, created_by=ctx.user_id))
    await audit(db, ctx, "workflow.create", "workflow", workflow.id, {"version": 1})
    await db.commit()
    return {"id": workflow.id, "name": workflow.name, "slug": workflow.slug, "latest_version": 1, "status": workflow.status}


@app.get("/api/workflows/{workflow_id}")
async def get_workflow(workflow_id: str, ctx: Context, db: AsyncSession = Depends(get_db)):
    workflow = await db.scalar(select(Workflow).where(Workflow.id == workflow_id, Workflow.tenant_id == ctx.tenant_id))
    if not workflow:
        raise HTTPException(status_code=404, detail="工作流不存在")
    version = await db.scalar(select(WorkflowVersion).where(WorkflowVersion.workflow_id == workflow.id, WorkflowVersion.version == workflow.latest_version))
    return {"id": workflow.id, "project_id": workflow.project_id, "name": workflow.name, "slug": workflow.slug, "description": workflow.description, "status": workflow.status, "latest_version": workflow.latest_version, "published_version": workflow.published_version, "graph": version.graph_json if version else {}, "input_schema": version.input_schema if version else {}}


@app.post("/api/workflows/{workflow_id}/validate")
async def validate_workflow(workflow_id: str, ctx: Context, db: AsyncSession = Depends(get_db)):
    workflow = await db.scalar(select(Workflow).where(Workflow.id == workflow_id, Workflow.tenant_id == ctx.tenant_id))
    if not workflow:
        raise HTTPException(status_code=404, detail="工作流不存在")
    version = await db.scalar(select(WorkflowVersion).where(WorkflowVersion.workflow_id == workflow.id, WorkflowVersion.version == workflow.latest_version))
    graph = version.graph_json if version else {}
    nodes = graph.get("nodes") or []
    edges = graph.get("edges") or []
    issues: list[dict[str, str]] = []
    ids = [str(node.get("id", "")) for node in nodes]
    if not nodes:
        issues.append({"level": "error", "code": "empty_graph", "message": "工作流没有节点"})
    if len(ids) != len(set(ids)):
        issues.append({"level": "error", "code": "duplicate_node_id", "message": "节点 ID 重复"})
    node_ids = set(ids)
    for edge in edges:
        if edge.get("source") not in node_ids or edge.get("target") not in node_ids:
            issues.append({"level": "error", "code": "dangling_edge", "message": f"边 {edge.get('id', '')} 引用了不存在的节点"})
    kinds = {str(node.get("data", {}).get("kind", node.get("type", ""))) for node in nodes}
    if "input" not in kinds:
        issues.append({"level": "error", "code": "missing_input", "message": "缺少输入节点"})
    if "output" not in kinds:
        issues.append({"level": "error", "code": "missing_output", "message": "缺少输出节点"})
    try:
        ordered_nodes(graph)
    except ValueError as exc:
        issues.append({"level": "error", "code": "invalid_graph", "message": str(exc)})
    return {"workflow_id": workflow.id, "version": workflow.latest_version, "valid": not any(item["level"] == "error" for item in issues), "issues": issues}


@app.get("/api/workflows/{workflow_id}/versions")
async def list_workflow_versions(workflow_id: str, ctx: Context, db: AsyncSession = Depends(get_db)):
    workflow = await db.scalar(select(Workflow).where(Workflow.id == workflow_id, Workflow.tenant_id == ctx.tenant_id))
    if not workflow:
        raise HTTPException(status_code=404, detail="工作流不存在")
    rows = (await db.execute(select(WorkflowVersion).where(WorkflowVersion.workflow_id == workflow.id, WorkflowVersion.tenant_id == ctx.tenant_id).order_by(desc(WorkflowVersion.version)))).scalars().all()
    return [{"id": row.id, "version": row.version, "input_schema": row.input_schema, "changelog": row.changelog, "created_by": row.created_by, "created_at": row.created_at, "published": row.version == workflow.published_version} for row in rows]


@app.get("/api/workflows/{workflow_id}/releases")
async def list_workflow_releases(workflow_id: str, ctx: Context, db: AsyncSession = Depends(get_db)):
    workflow = await db.scalar(select(Workflow).where(Workflow.id == workflow_id, Workflow.tenant_id == ctx.tenant_id))
    if not workflow:
        raise HTTPException(status_code=404, detail="工作流不存在")
    rows = (await db.execute(select(WorkflowRelease).where(WorkflowRelease.workflow_id == workflow.id, WorkflowRelease.tenant_id == ctx.tenant_id).order_by(desc(WorkflowRelease.created_at)).limit(100))).scalars().all()
    return [{"id": row.id, "version": row.version, "environment": row.environment, "traffic_percent": row.traffic_percent, "status": row.status, "released_by": row.released_by, "created_at": row.created_at} for row in rows]


@app.delete("/api/workflows/{workflow_id}")
async def delete_workflow(workflow_id: str, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    workflow = await db.scalar(select(Workflow).where(Workflow.id == workflow_id, Workflow.tenant_id == ctx.tenant_id))
    if not workflow:
        raise HTTPException(status_code=404, detail="工作流不存在")
    active_run = await db.scalar(select(WorkflowRun).where(WorkflowRun.workflow_id == workflow.id, WorkflowRun.tenant_id == ctx.tenant_id, WorkflowRun.status.in_(["queued", "running", "waiting_approval"])))
    if active_run:
        raise HTTPException(status_code=409, detail="工作流仍有运行中的任务，请先等待任务结束")
    workflow_id_value = workflow.id
    workflow_name = workflow.name
    await audit(db, ctx, "workflow.delete", "workflow", workflow_id_value, {"name": workflow_name})
    await db.execute(delete(ComputeUsage).where(ComputeUsage.run_id.in_(select(WorkflowRun.id).where(WorkflowRun.workflow_id == workflow_id_value, WorkflowRun.tenant_id == ctx.tenant_id))))
    await db.execute(delete(WorkflowRelease).where(WorkflowRelease.workflow_id == workflow_id_value, WorkflowRelease.tenant_id == ctx.tenant_id))
    await db.execute(delete(WorkflowSchedule).where(WorkflowSchedule.workflow_id == workflow_id_value, WorkflowSchedule.tenant_id == ctx.tenant_id))
    await db.execute(delete(WorkflowRun).where(WorkflowRun.workflow_id == workflow_id_value, WorkflowRun.tenant_id == ctx.tenant_id))
    await db.execute(delete(WorkflowVersion).where(WorkflowVersion.workflow_id == workflow_id_value, WorkflowVersion.tenant_id == ctx.tenant_id))
    await db.delete(workflow)
    await db.commit()
    return {"workflow_id": workflow_id_value, "deleted": True}


@app.post("/api/workflows/{workflow_id}/versions")
async def create_workflow_version(workflow_id: str, body: WorkflowVersionCreate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    workflow = await db.scalar(select(Workflow).where(Workflow.id == workflow_id, Workflow.tenant_id == ctx.tenant_id))
    if not workflow:
        raise HTTPException(status_code=404, detail="工作流不存在")
    try:
        ordered_nodes(body.graph)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    version_number = workflow.latest_version + 1
    db.add(WorkflowVersion(workflow_id=workflow.id, tenant_id=ctx.tenant_id, version=version_number, graph_json=body.graph, input_schema=body.input_schema, changelog=body.changelog, created_by=ctx.user_id))
    workflow.latest_version = version_number
    workflow.status = "draft"
    await audit(db, ctx, "workflow.version.create", "workflow", workflow.id, {"version": version_number})
    await db.commit()
    return {"workflow_id": workflow.id, "version": version_number, "status": workflow.status}


@app.post("/api/workflows/{workflow_id}/publish")
async def publish_workflow(workflow_id: str, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    workflow = await db.scalar(select(Workflow).where(Workflow.id == workflow_id, Workflow.tenant_id == ctx.tenant_id))
    if not workflow:
        raise HTTPException(status_code=404, detail="工作流不存在")
    workflow.published_version = workflow.latest_version
    workflow.status = "published"
    await audit(db, ctx, "workflow.publish", "workflow", workflow.id, {"version": workflow.published_version})
    await db.commit()
    return {"workflow_id": workflow.id, "published_version": workflow.published_version, "status": workflow.status}


@app.post("/api/workflows/{workflow_id}/release")
async def release_workflow(workflow_id: str, body: ReleaseRequest, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    workflow = await db.scalar(select(Workflow).where(Workflow.id == workflow_id, Workflow.tenant_id == ctx.tenant_id))
    if not workflow:
        raise HTTPException(status_code=404, detail="工作流不存在")
    version = body.version or workflow.published_version
    if not version:
        raise HTTPException(status_code=422, detail="请先发布工作流版本")
    if not await db.scalar(select(WorkflowVersion).where(WorkflowVersion.workflow_id == workflow.id, WorkflowVersion.version == version)):
        raise HTTPException(status_code=404, detail="版本不存在")
    release_status = "canary" if body.environment == "production" and body.traffic_percent < 100 else "active"
    if release_status == "active":
        active_releases = (await db.execute(select(WorkflowRelease).where(WorkflowRelease.workflow_id == workflow.id, WorkflowRelease.tenant_id == ctx.tenant_id, WorkflowRelease.environment == body.environment, WorkflowRelease.status == "active"))).scalars().all()
        for previous in active_releases:
            previous.status = "archived"
    else:
        canary_releases = (await db.execute(select(WorkflowRelease).where(WorkflowRelease.workflow_id == workflow.id, WorkflowRelease.tenant_id == ctx.tenant_id, WorkflowRelease.environment == body.environment, WorkflowRelease.status == "canary"))).scalars().all()
        for previous in canary_releases:
            previous.status = "archived"
    db.add(WorkflowRelease(tenant_id=ctx.tenant_id, workflow_id=workflow.id, version=version, environment=body.environment, traffic_percent=body.traffic_percent, status=release_status, released_by=ctx.user_id))
    await audit(db, ctx, "workflow.release", "workflow", workflow.id, {"version": version, "environment": body.environment, "traffic_percent": body.traffic_percent})
    await db.commit()
    return {"workflow_id": workflow.id, "version": version, "environment": body.environment, "traffic_percent": body.traffic_percent, "status": release_status}


@app.post("/api/workflows/{workflow_id}/rollback")
async def rollback_workflow(workflow_id: str, body: RollbackRequest, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    workflow = await db.scalar(select(Workflow).where(Workflow.id == workflow_id, Workflow.tenant_id == ctx.tenant_id))
    if not workflow:
        raise HTTPException(status_code=404, detail="工作流不存在")
    canaries = (await db.execute(select(WorkflowRelease).where(WorkflowRelease.workflow_id == workflow.id, WorkflowRelease.tenant_id == ctx.tenant_id, WorkflowRelease.environment == body.environment, WorkflowRelease.status == "canary"))).scalars().all()
    if not canaries:
        raise HTTPException(status_code=404, detail="没有可回滚的灰度版本")
    for canary in canaries:
        canary.status = "rolled_back"
    stable = await db.scalar(select(WorkflowRelease).where(WorkflowRelease.workflow_id == workflow.id, WorkflowRelease.tenant_id == ctx.tenant_id, WorkflowRelease.environment == body.environment, WorkflowRelease.status == "active").order_by(desc(WorkflowRelease.created_at)))
    await audit(db, ctx, "workflow.rollback", "workflow", workflow.id, {"environment": body.environment, "stable_version": stable.version if stable else None})
    await db.commit()
    return {"workflow_id": workflow.id, "environment": body.environment, "status": "rolled_back", "active_version": stable.version if stable else None}


async def select_release_version(db: AsyncSession, workflow: Workflow, environment: str, routing_key: str) -> int | None:
    releases = (await db.execute(select(WorkflowRelease).where(WorkflowRelease.workflow_id == workflow.id, WorkflowRelease.tenant_id == workflow.tenant_id, WorkflowRelease.environment == environment, WorkflowRelease.status.in_(["active", "canary"])).order_by(desc(WorkflowRelease.created_at)))).scalars().all()
    stable = next((item for item in releases if item.status == "active"), None)
    canary = next((item for item in releases if item.status == "canary"), None)
    if canary:
        bucket = int(hashlib.sha256(routing_key.encode()).hexdigest()[:8], 16) % 100 + 1
        if bucket <= canary.traffic_percent:
            return canary.version
    return stable.version if stable else (canary.version if canary else None)


@app.post("/api/workflows/{workflow_id}/runs")
async def run_workflow(workflow_id: str, body: WorkflowRunRequest, ctx: OperatorContext, db: AsyncSession = Depends(get_db)):
    workflow = await db.scalar(select(Workflow).where(Workflow.id == workflow_id, Workflow.tenant_id == ctx.tenant_id))
    if not workflow:
        raise HTTPException(status_code=404, detail="工作流不存在")
    version_number = body.version or await select_release_version(db, workflow, body.environment, f"{ctx.user_id}:{workflow.id}") or workflow.published_version
    if not version_number:
        raise HTTPException(status_code=422, detail="工作流尚未发布")
    version = await db.scalar(select(WorkflowVersion).where(WorkflowVersion.workflow_id == workflow.id, WorkflowVersion.version == version_number))
    if not version:
        raise HTTPException(status_code=404, detail="工作流版本不存在")
    schema_errors = []
    for field, rule in (version.input_schema or {}).items():
        if isinstance(rule, dict) and rule.get("required") and field not in body.input:
            schema_errors.append(f"缺少必填输入: {field}")
        if field not in body.input or not isinstance(rule, dict) or not rule.get("type"):
            continue
        expected = {"string": str, "str": str, "integer": int, "int": int, "number": (int, float), "boolean": bool, "bool": bool, "object": dict, "array": list}.get(str(rule["type"]))
        if expected and not isinstance(body.input[field], expected):
            schema_errors.append(f"输入 {field} 类型应为 {rule['type']}")
    if schema_errors:
        raise HTTPException(status_code=422, detail="；".join(schema_errors))
    quota = await db.scalar(select(ComputeQuota).where(ComputeQuota.tenant_id == ctx.tenant_id))
    if quota and quota.enabled:
        running = await db.scalar(select(func.count()).select_from(WorkflowRun).where(WorkflowRun.tenant_id == ctx.tenant_id, WorkflowRun.status.in_(["queued", "running"]))) or 0
        if int(running) >= quota.concurrent_runs:
            raise HTTPException(status_code=429, detail="已达到租户并发运行配额")
    run = WorkflowRun(tenant_id=ctx.tenant_id, project_id=workflow.project_id, workflow_id=workflow.id, workflow_version=version_number, input_json=body.input, trace_id="", created_by=ctx.user_id)
    db.add(run)
    await audit(db, ctx, "workflow.run.create", "workflow_run", run.id, {"workflow_id": workflow.id, "version": version_number})
    await db.commit()
    await enqueue_run(run.id)
    return {"run_id": run.id, "status": run.status, "workflow_id": workflow.id, "version": version_number}


@app.post("/api/v1/workflows/{workflow_id}/run")
async def public_run_workflow(workflow_id: str, body: WorkflowRunRequest, ctx: OperatorContext, db: AsyncSession = Depends(get_db)):
    """API Key/JWT 对外 Workflow 调用入口。"""
    require_api_scope(ctx, "workflow:run")
    return await run_workflow(workflow_id, body, ctx, db)


def workflow_webhook_token(workflow_id: str) -> str:
    return hashlib.sha256(f"{get_settings().app_secret_key}:{workflow_id}".encode()).hexdigest()


def require_api_scope(ctx: RequestContext, scope: str) -> None:
    """JWT 管理调用不受 API Key scope 限制；API Key 必须显式声明调用范围。"""
    api_scopes = {item[4:] for item in ctx.custom_permissions if item.startswith("api:")}
    if api_scopes and scope not in api_scopes:
        raise HTTPException(status_code=403, detail=f"API Key 缺少权限: {scope}")


@app.get("/api/workflows/{workflow_id}/webhook-token")
async def get_workflow_webhook_token(workflow_id: str, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    workflow = await db.scalar(select(Workflow).where(Workflow.id == workflow_id, Workflow.tenant_id == ctx.tenant_id))
    if not workflow:
        raise HTTPException(status_code=404, detail="工作流不存在")
    return {"workflow_id": workflow.id, "token": workflow_webhook_token(workflow.id)}


@app.post("/api/workflows/{workflow_id}/webhook")
async def trigger_workflow_webhook(workflow_id: str, body: WorkflowWebhookRequest, x_workflow_token: str | None = Header(default=None), db: AsyncSession = Depends(get_admin_db)):
    workflow = await db.scalar(select(Workflow).where(Workflow.id == workflow_id, Workflow.status == "published"))
    if not workflow or not x_workflow_token or not hmac.compare_digest(x_workflow_token, workflow_webhook_token(workflow.id)):
        raise HTTPException(status_code=404, detail="工作流不存在或 Webhook 凭证无效")
    version_number = body.version or await select_release_version(db, workflow, "production", f"webhook:{workflow.id}") or workflow.published_version
    if not version_number:
        raise HTTPException(status_code=422, detail="工作流尚未发布")
    run = WorkflowRun(tenant_id=workflow.tenant_id, project_id=workflow.project_id, workflow_id=workflow.id, workflow_version=version_number, input_json=body.input, trace_id="", created_by=workflow.created_by)
    db.add(run)
    await db.commit()
    await enqueue_run(run.id)
    return {"run_id": run.id, "status": run.status, "workflow_id": workflow.id, "version": version_number}


@app.get("/api/workflows/{workflow_id}/schedules")
async def list_workflow_schedules(workflow_id: str, ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(WorkflowSchedule).where(WorkflowSchedule.workflow_id == workflow_id, WorkflowSchedule.tenant_id == ctx.tenant_id).order_by(WorkflowSchedule.created_at.desc()))).scalars().all()
    return [{"id": row.id, "workflow_id": row.workflow_id, "environment": row.environment, "interval_seconds": row.interval_seconds, "next_run_at": row.next_run_at, "last_run_at": row.last_run_at, "input": row.input_json, "active": row.active} for row in rows]


@app.post("/api/workflows/{workflow_id}/schedules")
async def create_workflow_schedule(workflow_id: str, body: ScheduleCreate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    workflow = await db.scalar(select(Workflow).where(Workflow.id == workflow_id, Workflow.tenant_id == ctx.tenant_id))
    if not workflow:
        raise HTTPException(status_code=404, detail="工作流不存在")
    row = WorkflowSchedule(tenant_id=ctx.tenant_id, workflow_id=workflow_id, environment=body.environment, interval_seconds=body.interval_seconds, next_run_at=utcnow() + timedelta(seconds=body.interval_seconds), input_json=body.input, active=body.active, created_by=ctx.user_id)
    db.add(row)
    await audit(db, ctx, "workflow.schedule.create", "workflow_schedule", row.id, {"interval_seconds": row.interval_seconds})
    await db.commit()
    return {"id": row.id, "workflow_id": row.workflow_id, "interval_seconds": row.interval_seconds, "next_run_at": row.next_run_at, "active": row.active}


@app.patch("/api/workflow-schedules/{schedule_id}")
async def update_workflow_schedule(schedule_id: str, body: ScheduleUpdate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(WorkflowSchedule).where(WorkflowSchedule.id == schedule_id, WorkflowSchedule.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="定时任务不存在")
    if body.interval_seconds is not None:
        row.interval_seconds = body.interval_seconds
        row.next_run_at = utcnow() + timedelta(seconds=body.interval_seconds)
    if body.active is not None:
        row.active = body.active
    if body.input is not None:
        row.input_json = body.input
    await db.commit()
    return {"id": row.id, "interval_seconds": row.interval_seconds, "next_run_at": row.next_run_at, "active": row.active, "input": row.input_json}


@app.delete("/api/workflow-schedules/{schedule_id}")
async def delete_workflow_schedule(schedule_id: str, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(WorkflowSchedule).where(WorkflowSchedule.id == schedule_id, WorkflowSchedule.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="定时任务不存在")
    await db.delete(row)
    await db.commit()
    return {"id": schedule_id, "deleted": True}


@app.get("/api/runs")
async def list_runs(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(WorkflowRun).where(WorkflowRun.tenant_id == ctx.tenant_id).order_by(desc(WorkflowRun.created_at)).limit(100))).scalars().all()
    return [{"id": row.id, "workflow_id": row.workflow_id, "status": row.status, "version": row.workflow_version, "trace_id": row.trace_id, "error": row.error, "created_at": row.created_at, "finished_at": row.finished_at} for row in rows]


@app.get("/api/runs/{run_id}")
async def get_run(run_id: str, ctx: Context, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(WorkflowRun).where(WorkflowRun.id == run_id, WorkflowRun.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="运行记录不存在")
    return {"id": row.id, "workflow_id": row.workflow_id, "status": row.status, "version": row.workflow_version, "input": row.input_json, "state": row.state_json, "output": row.output_json, "error": row.error, "trace_id": row.trace_id, "created_at": row.created_at, "started_at": row.started_at, "finished_at": row.finished_at}


@app.post("/api/runs/{run_id}/approval")
async def approve_run(run_id: str, body: ApprovalRequest, ctx: OperatorContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(WorkflowRun).where(WorkflowRun.id == run_id, WorkflowRun.tenant_id == ctx.tenant_id, WorkflowRun.status == "waiting_approval"))
    if not row:
        raise HTTPException(status_code=404, detail="没有等待审批的运行记录")
    state = dict(row.state_json or {})
    waiting_node = state.get("waiting_node_id")
    nodes = dict(state.get("nodes", {}))
    nodes[str(waiting_node)] = {"approved": body.approved, "comment": body.comment}
    state["nodes"] = nodes
    row.state_json = state
    row.status = "queued" if body.approved else "failed"
    row.error = "审批拒绝" if not body.approved else ""
    if not body.approved:
        row.finished_at = utcnow()
    await audit(db, ctx, "workflow.run.approval", "workflow_run", row.id, {"approved": body.approved, "comment": body.comment})
    await db.commit()
    if body.approved:
        await enqueue_run(row.id)
    return {"run_id": row.id, "status": row.status, "approved": body.approved}


@app.get("/api/knowledge-bases")
async def list_kbs(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(KnowledgeBase).where(KnowledgeBase.tenant_id == ctx.tenant_id).order_by(KnowledgeBase.name))).scalars().all()
    return [{"id": row.id, "project_id": row.project_id, "name": row.name, "description": row.description, "created_at": row.created_at} for row in rows]


async def require_kb_access(kb_id: str, ctx: RequestContext, db: AsyncSession, write: bool = False) -> None:
    total = await db.scalar(select(func.count()).select_from(KnowledgeBaseAccess).where(KnowledgeBaseAccess.knowledge_base_id == kb_id, KnowledgeBaseAccess.tenant_id == ctx.tenant_id)) or 0
    if not total:
        return
    access = await db.scalar(select(KnowledgeBaseAccess).where(KnowledgeBaseAccess.knowledge_base_id == kb_id, KnowledgeBaseAccess.tenant_id == ctx.tenant_id, KnowledgeBaseAccess.user_id == ctx.user_id))
    if not access or (write and access.access_level not in {"owner", "editor"}):
        raise HTTPException(status_code=403, detail="没有该知识库的访问权限")


@app.post("/api/knowledge-bases")
async def create_kb(body: KnowledgeBaseCreate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    project = await db.scalar(select(Project).where(Project.id == body.project_id, Project.tenant_id == ctx.tenant_id))
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    row = KnowledgeBase(tenant_id=ctx.tenant_id, project_id=project.id, name=body.name, description=body.description, created_by=ctx.user_id)
    db.add(row)
    await db.flush()
    db.add(KnowledgeBaseAccess(tenant_id=ctx.tenant_id, knowledge_base_id=row.id, user_id=ctx.user_id, access_level="owner"))
    await audit(db, ctx, "knowledge_base.create", "knowledge_base", row.id)
    await db.commit()
    return {"id": row.id, "name": row.name, "project_id": row.project_id}


@app.get("/api/knowledge-bases/{kb_id}/access")
async def list_kb_access(kb_id: str, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    kb = await db.scalar(select(KnowledgeBase).where(KnowledgeBase.id == kb_id, KnowledgeBase.tenant_id == ctx.tenant_id))
    if not kb:
        raise HTTPException(status_code=404, detail="知识库不存在")
    result = await db.execute(select(KnowledgeBaseAccess, User).join(User, User.id == KnowledgeBaseAccess.user_id).where(KnowledgeBaseAccess.knowledge_base_id == kb_id, KnowledgeBaseAccess.tenant_id == ctx.tenant_id))
    return [{"user_id": access.user_id, "email": user.email, "display_name": user.display_name, "access_level": access.access_level} for access, user in result.all()]


@app.post("/api/knowledge-bases/{kb_id}/access")
async def grant_kb_access(kb_id: str, body: KnowledgeBaseAccessCreate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    kb = await db.scalar(select(KnowledgeBase).where(KnowledgeBase.id == kb_id, KnowledgeBase.tenant_id == ctx.tenant_id))
    user = await db.scalar(select(User).where(User.id == body.user_id))
    if not kb or not user:
        raise HTTPException(status_code=404, detail="知识库或用户不存在")
    access = await db.scalar(select(KnowledgeBaseAccess).where(KnowledgeBaseAccess.knowledge_base_id == kb_id, KnowledgeBaseAccess.user_id == user.id))
    if not access:
        access = KnowledgeBaseAccess(tenant_id=ctx.tenant_id, knowledge_base_id=kb_id, user_id=user.id, access_level=body.access_level)
        db.add(access)
    else:
        access.access_level = body.access_level
    await audit(db, ctx, "knowledge.access.grant", "knowledge_base", kb_id, {"user_id": user.id, "access_level": body.access_level})
    await db.commit()
    return {"knowledge_base_id": kb_id, "user_id": user.id, "access_level": access.access_level}


@app.post("/api/knowledge-bases/{kb_id}/documents")
async def upload_document(kb_id: str, ctx: EditorContext, file: UploadFile = File(...), db: AsyncSession = Depends(get_db)):
    kb = await db.scalar(select(KnowledgeBase).where(KnowledgeBase.id == kb_id, KnowledgeBase.tenant_id == ctx.tenant_id))
    if not kb:
        raise HTTPException(status_code=404, detail="知识库不存在")
    await require_kb_access(kb_id, ctx, db, write=True)
    raw = await file.read()
    filename = (file.filename or "").lower()
    if filename.endswith(".pdf"):
        try:
            reader = PdfReader(io.BytesIO(raw))
            content = "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"PDF 解析失败: {exc}") from exc
    elif filename.endswith(".docx"):
        try:
            docx = DocxDocument(io.BytesIO(raw))
            content = "\n".join(paragraph.text for paragraph in docx.paragraphs)
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"DOCX 解析失败: {exc}") from exc
    elif filename.endswith((".xlsx", ".xlsm")):
        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            content = "\n".join("\t".join("" if value is None else str(value) for value in row) for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True))
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"Excel 解析失败: {exc}") from exc
    else:
        content = raw.decode("utf-8", errors="replace")
    if not content.strip():
        raise HTTPException(status_code=422, detail="文档没有可索引的文本内容")
    document = Document(tenant_id=ctx.tenant_id, kb_id=kb.id, title=file.filename or "uploaded-document", source="upload", content=content, checksum=checksum(raw), status="indexing", created_by=ctx.user_id)
    db.add(document)
    await db.flush()
    chunks = []
    for index, piece in enumerate(chunk_text(content)):
        chunk = DocumentChunk(tenant_id=ctx.tenant_id, kb_id=kb.id, document_id=document.id, chunk_index=index, content=piece)
        db.add(chunk)
        chunks.append(chunk)
    await db.flush()
    await audit(db, ctx, "knowledge.document.upload", "document", document.id, {"kb_id": kb.id, "filename": file.filename})
    await db.commit()
    await enqueue_index(document.id)
    return {"id": document.id, "kb_id": kb.id, "title": document.title, "status": document.status, "checksum": document.checksum, "vector_indexed": 0, "index_queued": True}


@app.get("/api/knowledge-bases/{kb_id}/documents")
async def list_documents(kb_id: str, ctx: Context, db: AsyncSession = Depends(get_db)):
    await require_kb_access(kb_id, ctx, db)
    rows = (await db.execute(select(Document).where(Document.kb_id == kb_id, Document.tenant_id == ctx.tenant_id).order_by(Document.created_at.desc()))).scalars().all()
    return [{"id": row.id, "kb_id": row.kb_id, "title": row.title, "source": row.source, "checksum": row.checksum, "status": row.status, "created_at": row.created_at} for row in rows]


@app.post("/api/documents/{document_id}/reindex")
async def reindex_document(document_id: str, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    document = await db.scalar(select(Document).where(Document.id == document_id, Document.tenant_id == ctx.tenant_id))
    if not document:
        raise HTTPException(status_code=404, detail="文档不存在")
    await require_kb_access(document.kb_id, ctx, db, write=True)
    document.status = "indexing"
    await db.commit()
    await enqueue_index(document.id)
    return {"id": document.id, "status": document.status, "index_queued": True}


@app.delete("/api/documents/{document_id}")
async def delete_document(document_id: str, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    document = await db.scalar(select(Document).where(Document.id == document_id, Document.tenant_id == ctx.tenant_id))
    if not document:
        raise HTTPException(status_code=404, detail="文档不存在")
    await require_kb_access(document.kb_id, ctx, db, write=True)
    await db.execute(delete(DocumentChunk).where(DocumentChunk.document_id == document.id, DocumentChunk.tenant_id == ctx.tenant_id))
    await db.delete(document)
    await audit(db, ctx, "knowledge.document.delete", "document", document_id)
    await db.commit()
    try:
        from .services.vector import delete_document_vectors
        await delete_document_vectors(ctx.tenant_id, document.kb_id, document_id)
    except Exception:
        logger.warning("failed to remove document vectors", exc_info=True)
    return {"id": document_id, "deleted": True}


@app.get("/api/knowledge-bases/{kb_id}/search")
async def query_kb(kb_id: str, q: str, ctx: Context, mode: str = "hybrid", db: AsyncSession = Depends(get_db)):
    kb = await db.scalar(select(KnowledgeBase).where(KnowledgeBase.id == kb_id, KnowledgeBase.tenant_id == ctx.tenant_id))
    if not kb:
        raise HTTPException(status_code=404, detail="知识库不存在")
    await require_kb_access(kb_id, ctx, db)
    if mode not in {"hybrid", "keyword"}:
        raise HTTPException(status_code=422, detail="检索模式必须是 hybrid 或 keyword")
    return {"query": q, "mode": mode, "results": await search_chunks(db, ctx.tenant_id, kb.id, q, mode=mode)}


@app.get("/api/knowledge-bases/{kb_id}/assets")
async def list_kb_assets(kb_id: str, ctx: Context, asset_type: str | None = None, db: AsyncSession = Depends(get_db)):
    kb = await db.scalar(select(KnowledgeBase).where(KnowledgeBase.id == kb_id, KnowledgeBase.tenant_id == ctx.tenant_id))
    if not kb:
        raise HTTPException(status_code=404, detail="知识库不存在")
    await require_kb_access(kb_id, ctx, db)
    stmt = select(KnowledgeAsset).where(KnowledgeAsset.kb_id == kb_id, KnowledgeAsset.tenant_id == ctx.tenant_id)
    if asset_type:
        stmt = stmt.where(KnowledgeAsset.asset_type == asset_type)
    rows = (await db.execute(stmt.order_by(desc(KnowledgeAsset.updated_at)))).scalars().all()
    return [{"id": row.id, "kb_id": row.kb_id, "asset_type": row.asset_type, "name": row.name, "description": row.description, "content": row.content_json, "status": row.status, "created_at": row.created_at, "updated_at": row.updated_at} for row in rows]


@app.post("/api/knowledge-bases/{kb_id}/assets")
async def create_kb_asset(kb_id: str, body: KnowledgeAssetCreate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    kb = await db.scalar(select(KnowledgeBase).where(KnowledgeBase.id == kb_id, KnowledgeBase.tenant_id == ctx.tenant_id))
    if not kb:
        raise HTTPException(status_code=404, detail="知识库不存在")
    await require_kb_access(kb_id, ctx, db, write=True)
    row = KnowledgeAsset(tenant_id=ctx.tenant_id, kb_id=kb_id, asset_type=body.asset_type, name=body.name, description=body.description, content_json=body.content, status=body.status, created_by=ctx.user_id)
    db.add(row)
    try:
        await db.flush()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(status_code=409, detail="同一知识库中已存在同名资源") from exc
    await audit(db, ctx, "knowledge.asset.create", "knowledge_asset", row.id, {"kb_id": kb_id, "asset_type": body.asset_type})
    await db.commit()
    return {"id": row.id, "kb_id": row.kb_id, "asset_type": row.asset_type, "name": row.name, "status": row.status}


@app.patch("/api/knowledge-assets/{asset_id}")
async def update_kb_asset(asset_id: str, body: KnowledgeAssetCreate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(KnowledgeAsset).where(KnowledgeAsset.id == asset_id, KnowledgeAsset.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="知识资源不存在")
    await require_kb_access(row.kb_id, ctx, db, write=True)
    row.name, row.description, row.content_json, row.status = body.name, body.description, body.content, body.status
    await audit(db, ctx, "knowledge.asset.update", "knowledge_asset", row.id, {"status": row.status})
    await db.commit()
    return {"id": row.id, "name": row.name, "status": row.status}


@app.delete("/api/knowledge-assets/{asset_id}")
async def delete_kb_asset(asset_id: str, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(KnowledgeAsset).where(KnowledgeAsset.id == asset_id, KnowledgeAsset.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="知识资源不存在")
    await require_kb_access(row.kb_id, ctx, db, write=True)
    await db.delete(row)
    await audit(db, ctx, "knowledge.asset.delete", "knowledge_asset", asset_id)
    await db.commit()
    return {"id": asset_id, "deleted": True}


@app.get("/api/skills")
async def list_skills(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(Skill).where(Skill.tenant_id == ctx.tenant_id).order_by(Skill.name, Skill.version))).scalars().all()
    return [{"id": row.id, "project_id": row.project_id, "name": row.name, "version": row.version, "description": row.description, "enabled": row.enabled, "manifest": row.manifest_json} for row in rows]


@app.post("/api/skills")
async def create_skill(body: SkillCreate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    project = await db.scalar(select(Project).where(Project.id == body.project_id, Project.tenant_id == ctx.tenant_id))
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    row = Skill(tenant_id=ctx.tenant_id, project_id=project.id, name=body.name, version=body.version, description=body.description, instructions=body.instructions, manifest_json=body.manifest, created_by=ctx.user_id)
    db.add(row)
    await db.flush()
    await audit(db, ctx, "skill.create", "skill", row.id, {"name": row.name, "version": row.version})
    await db.commit()
    return {"id": row.id, "name": row.name, "version": row.version, "enabled": row.enabled}


@app.patch("/api/skills/{skill_id}")
async def update_skill(skill_id: str, body: EnabledUpdate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(Skill).where(Skill.id == skill_id, Skill.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="技能不存在")
    row.enabled = body.enabled
    await audit(db, ctx, "skill.update", "skill", row.id, {"enabled": row.enabled})
    await db.commit()
    return {"id": row.id, "enabled": row.enabled}


@app.post("/api/skills/{skill_id}/run")
async def run_skill(skill_id: str, body: dict[str, Any], ctx: OperatorContext, db: AsyncSession = Depends(get_db)):
    skill = await db.scalar(select(Skill).where(Skill.id == skill_id, Skill.tenant_id == ctx.tenant_id, Skill.enabled.is_(True)))
    if not skill:
        raise HTTPException(status_code=404, detail="技能不存在或已停用")
    manifest = skill.manifest_json or {}
    provider_stmt = select(ModelProvider).where(ModelProvider.tenant_id == ctx.tenant_id, ModelProvider.enabled.is_(True))
    provider_id = manifest.get("provider_id")
    if provider_id:
        provider_stmt = provider_stmt.where(ModelProvider.id == str(provider_id))
    provider = (await db.execute(provider_stmt)).scalars().first()
    if not provider:
        return {"skill_id": skill.id, "mode": "instruction", "output": skill.instructions, "input": body}
    messages = [{"role": "system", "content": skill.instructions}, {"role": "user", "content": json.dumps(body, ensure_ascii=False)}]
    result = await call_chat(db, provider, messages, tenant_id=ctx.tenant_id)
    await audit(db, ctx, "skill.run", "skill", skill.id, {"provider": provider.name})
    await db.commit()
    return {"skill_id": skill.id, "mode": "llm", "output": result.get("text", ""), "usage": result.get("usage", {})}


@app.get("/api/plugins")
async def list_plugins(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(Plugin).where(Plugin.tenant_id == ctx.tenant_id).order_by(Plugin.name))).scalars().all()
    return [{"id": row.id, "name": row.name, "kind": row.kind, "endpoint": row.endpoint, "manifest": row.manifest_json, "enabled": row.enabled, "approval_status": row.approval_status, "approved_at": row.approved_at} for row in rows]


@app.post("/api/plugins")
async def create_plugin(body: PluginCreate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    validate_extension_manifest(body.manifest)
    row = Plugin(tenant_id=ctx.tenant_id, name=body.name, kind=body.kind, endpoint=body.endpoint, manifest_json=body.manifest, secret_encrypted=encrypt_secret(body.secret), created_by=ctx.user_id)
    db.add(row)
    try:
        await db.flush()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(status_code=409, detail="插件名称已存在") from exc
    await audit(db, ctx, "plugin.create", "plugin", row.id, {"name": row.name})
    await db.commit()
    return {"id": row.id, "name": row.name, "kind": row.kind, "endpoint": row.endpoint, "enabled": row.enabled, "approval_status": row.approval_status}


@app.patch("/api/plugins/{plugin_id}")
async def update_plugin(plugin_id: str, body: EnabledUpdate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(Plugin).where(Plugin.id == plugin_id, Plugin.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="插件不存在")
    row.enabled = body.enabled
    await audit(db, ctx, "plugin.update", "plugin", row.id, {"enabled": row.enabled})
    await db.commit()
    return {"id": row.id, "enabled": row.enabled}


@app.post("/api/plugins/{plugin_id}/approve")
async def approve_plugin(plugin_id: str, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(Plugin).where(Plugin.id == plugin_id, Plugin.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="插件不存在")
    row.approval_status, row.approved_by, row.approved_at = "approved", ctx.user_id, utcnow()
    await audit(db, ctx, "plugin.approve", "plugin", row.id)
    await db.commit()
    return {"id": row.id, "approval_status": row.approval_status, "approved_at": row.approved_at}


@app.post("/api/plugins/{plugin_id}/reject")
async def reject_plugin(plugin_id: str, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(Plugin).where(Plugin.id == plugin_id, Plugin.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="插件不存在")
    row.approval_status, row.approved_by, row.approved_at = "rejected", ctx.user_id, utcnow()
    row.enabled = False
    await audit(db, ctx, "plugin.reject", "plugin", row.id)
    await db.commit()
    return {"id": row.id, "approval_status": row.approval_status, "enabled": row.enabled}


@app.post("/api/plugins/{plugin_id}/test")
async def test_plugin(plugin_id: str, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(Plugin).where(Plugin.id == plugin_id, Plugin.tenant_id == ctx.tenant_id, Plugin.enabled.is_(True)))
    if not row:
        raise HTTPException(status_code=404, detail="插件不存在或已停用")
    if row.approval_status != "approved":
        raise HTTPException(status_code=403, detail="插件必须先通过审批")
    from urllib.parse import urlparse
    host = (urlparse(row.endpoint).hostname or "").lower()
    if not host or host not in get_settings().allowed_hosts:
        raise HTTPException(status_code=422, detail="插件地址不在允许的主机列表中")
    headers = {"X-AI-Tenant-ID": ctx.tenant_id}
    if row.secret_encrypted:
        headers["X-AI-Plugin-Secret"] = decrypt_secret(row.secret_encrypted)
    async with httpx.AsyncClient(timeout=10) as client:
        response = await client.get(row.endpoint, headers=headers)
    await audit(db, ctx, "plugin.test", "plugin", row.id, {"status_code": response.status_code})
    await db.commit()
    return {"id": row.id, "status_code": response.status_code, "ok": response.is_success}


@app.get("/api/mcp-servers")
async def list_mcp_servers(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(MCPServer).where(MCPServer.tenant_id == ctx.tenant_id).order_by(MCPServer.name))).scalars().all()
    return [{"id": row.id, "name": row.name, "endpoint": row.endpoint, "transport": row.transport, "manifest": row.manifest_json, "enabled": row.enabled, "has_secret": bool(row.secret_encrypted)} for row in rows]


@app.post("/api/mcp-servers")
async def create_mcp_server(body: MCPServerCreate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    validate_extension_manifest(body.manifest)
    row = MCPServer(tenant_id=ctx.tenant_id, name=body.name, endpoint=body.endpoint, transport=body.transport, manifest_json=body.manifest, secret_encrypted=encrypt_secret(body.secret), created_by=ctx.user_id)
    db.add(row)
    try:
        await db.flush()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(status_code=409, detail="MCP 服务名称已存在") from exc
    await audit(db, ctx, "mcp_server.create", "mcp_server", row.id, {"transport": row.transport})
    await db.commit()
    return {"id": row.id, "name": row.name, "endpoint": row.endpoint, "transport": row.transport, "enabled": row.enabled}


@app.patch("/api/mcp-servers/{server_id}")
async def update_mcp_server(server_id: str, body: EnabledUpdate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(MCPServer).where(MCPServer.id == server_id, MCPServer.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="MCP 服务不存在")
    row.enabled = body.enabled
    await audit(db, ctx, "mcp_server.update", "mcp_server", row.id, {"enabled": row.enabled})
    await db.commit()
    return {"id": row.id, "enabled": row.enabled}


@app.post("/api/mcp-servers/{server_id}/test")
async def test_mcp_server(server_id: str, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(MCPServer).where(MCPServer.id == server_id, MCPServer.tenant_id == ctx.tenant_id, MCPServer.enabled.is_(True)))
    if not row:
        raise HTTPException(status_code=404, detail="MCP 服务不存在或已停用")
    from urllib.parse import urlparse
    host = (urlparse(row.endpoint).hostname or "").lower()
    if row.transport != "stdio" and (not host or host not in get_settings().allowed_hosts):
        raise HTTPException(status_code=422, detail="MCP 服务地址不在允许的主机列表中")
    headers = {"X-AI-Tenant-ID": ctx.tenant_id}
    if row.secret_encrypted:
        headers["Authorization"] = f"Bearer {decrypt_secret(row.secret_encrypted)}"
    status_code = 0
    detail = "stdio 服务已登记"
    if row.transport != "stdio":
        async with httpx.AsyncClient(timeout=10) as client:
            response = await client.get(row.endpoint, headers=headers)
        status_code = response.status_code
        detail = response.text[:500]
    await audit(db, ctx, "mcp_server.test", "mcp_server", row.id, {"status_code": status_code})
    await db.commit()
    return {"id": row.id, "status_code": status_code, "ok": row.transport == "stdio" or 200 <= status_code < 500, "detail": detail}


@app.get("/api/audit")
async def list_audit(ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(AuditEvent).where(AuditEvent.tenant_id == ctx.tenant_id).order_by(desc(AuditEvent.created_at)).limit(200))).scalars().all()
    return [{"id": row.id, "action": row.action, "resource_type": row.resource_type, "resource_id": row.resource_id, "actor_id": row.actor_id, "detail": row.detail_json, "created_at": row.created_at} for row in rows]


@app.get("/api/audit/export", response_class=PlainTextResponse)
async def export_audit(ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(AuditEvent).where(AuditEvent.tenant_id == ctx.tenant_id).order_by(desc(AuditEvent.created_at)).limit(5000))).scalars().all()
    lines = ["id,action,resource_type,resource_id,actor_id,created_at,detail"]
    for row in rows:
        values = [row.id, row.action, row.resource_type, row.resource_id, row.actor_id or "", row.created_at.isoformat(), json.dumps(row.detail_json or {}, ensure_ascii=False)]
        lines.append(",".join('"' + str(value).replace('"', '""') + '"' for value in values))
    return PlainTextResponse("\n".join(lines), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": "attachment; filename=audit-events.csv"})


@app.get("/api/observability/summary")
async def observability_summary(ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    active_runs = await db.scalar(select(func.count()).select_from(WorkflowRun).where(WorkflowRun.tenant_id == ctx.tenant_id, WorkflowRun.status.in_(["queued", "running", "waiting_approval"]))) or 0
    failed_runs = await db.scalar(select(func.count()).select_from(WorkflowRun).where(WorkflowRun.tenant_id == ctx.tenant_id, WorkflowRun.status == "failed")) or 0
    total_runs = await db.scalar(select(func.count()).select_from(WorkflowRun).where(WorkflowRun.tenant_id == ctx.tenant_id)) or 0
    usage = (await db.execute(select(func.coalesce(func.sum(ComputeUsage.cpu_milliseconds), 0), func.coalesce(func.sum(ComputeUsage.gpu_seconds), 0), func.coalesce(func.sum(ComputeUsage.tokens), 0)).where(ComputeUsage.tenant_id == ctx.tenant_id))).one()
    return {"active_runs": int(active_runs), "failed_runs": int(failed_runs), "total_runs": int(total_runs), "failure_rate": round(float(failed_runs) / total_runs, 4) if total_runs else 0, "usage": {"cpu_milliseconds": int(usage[0]), "gpu_seconds": int(usage[1]), "tokens": int(usage[2])}, "security": secret_backend_status()}


# Product surface APIs. These resources back the workspace, Agent editor, model
# experience, marketplace, publishing and platform-admin screens.

@app.get("/api/workspaces")
async def list_workspaces(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(Workspace).where(Workspace.tenant_id == ctx.tenant_id).order_by(Workspace.name))).scalars().all()
    return [{"id": row.id, "name": row.name, "slug": row.slug, "type": row.workspace_type, "organization_id": row.organization_id, "owner_id": row.owner_id, "description": row.description, "quota": row.quota_json} for row in rows]


@app.post("/api/workspaces")
async def create_workspace(body: WorkspaceCreate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = Workspace(tenant_id=ctx.tenant_id, organization_id=body.organization_id, owner_id=ctx.user_id, name=body.name, slug=body.slug, workspace_type=body.workspace_type, description=body.description)
    db.add(row)
    await db.flush()
    db.add(WorkspaceMember(tenant_id=ctx.tenant_id, workspace_id=row.id, user_id=ctx.user_id, role="owner"))
    await audit(db, ctx, "workspace.create", "workspace", row.id, {"name": row.name})
    await db.commit()
    return {"id": row.id, "name": row.name, "slug": row.slug, "type": row.workspace_type}


@app.get("/api/organizations")
async def list_organizations(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(Organization).where(Organization.tenant_id == ctx.tenant_id).order_by(Organization.name))).scalars().all()
    return [{"id": row.id, "name": row.name, "code": row.code, "parent_id": row.parent_id} for row in rows]


@app.post("/api/organizations")
async def create_organization(body: OrganizationCreate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = Organization(tenant_id=ctx.tenant_id, parent_id=body.parent_id, name=body.name, code=body.code)
    db.add(row)
    await db.flush()
    await audit(db, ctx, "organization.create", "organization", row.id, {"name": row.name})
    await db.commit()
    return {"id": row.id, "name": row.name, "code": row.code, "parent_id": row.parent_id}


async def get_or_create_agent_config(db: AsyncSession, ctx: RequestContext, agent: Agent) -> AgentConfig:
    config = await db.scalar(select(AgentConfig).where(AgentConfig.agent_id == agent.id, AgentConfig.tenant_id == ctx.tenant_id))
    if config:
        return config
    workspace = await db.scalar(select(Workspace).where(Workspace.tenant_id == ctx.tenant_id).order_by(Workspace.created_at))
    config = AgentConfig(tenant_id=ctx.tenant_id, agent_id=agent.id, workspace_id=workspace.id if workspace else None, prompt_template=agent.system_prompt, capabilities_json={"knowledge_base": bool(agent.knowledge_base_id), "tools": []})
    db.add(config)
    await db.flush()
    return config


@app.get("/api/agents/{agent_id}/config")
async def get_agent_config(agent_id: str, ctx: Context, db: AsyncSession = Depends(get_db)):
    agent = await db.scalar(select(Agent).where(Agent.id == agent_id, Agent.tenant_id == ctx.tenant_id))
    if not agent:
        raise HTTPException(status_code=404, detail="智能体不存在")
    config = await get_or_create_agent_config(db, ctx, agent)
    await db.commit()
    return {"agent_id": agent.id, "name": agent.name, "description": agent.description, "provider_id": agent.provider_id, "knowledge_base_id": agent.knowledge_base_id, "agent_type": config.agent_type, "workspace_id": config.workspace_id, "logo_url": config.logo_url, "prompt_mode": config.prompt_mode, "prompt_template": config.prompt_template, "variables": config.variables_json, "memory": config.memory_json, "capabilities": config.capabilities_json, "experience": config.experience_json, "safety": config.safety_json, "publish_status": config.publish_status, "published_version": config.published_version}


@app.patch("/api/agents/{agent_id}/config")
async def update_agent_config(agent_id: str, body: AgentConfigUpdate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    agent = await db.scalar(select(Agent).where(Agent.id == agent_id, Agent.tenant_id == ctx.tenant_id))
    if not agent:
        raise HTTPException(status_code=404, detail="智能体不存在")
    if body.workspace_id and not await db.scalar(select(Workspace).where(Workspace.id == body.workspace_id, Workspace.tenant_id == ctx.tenant_id)):
        raise HTTPException(status_code=404, detail="工作空间不存在")
    config = await get_or_create_agent_config(db, ctx, agent)
    config.agent_type = body.agent_type
    config.workspace_id = body.workspace_id
    config.logo_url = body.logo_url
    config.prompt_mode = body.prompt_mode
    config.prompt_template = body.prompt_template
    config.variables_json = body.variables
    config.memory_json = body.memory
    config.capabilities_json = body.capabilities
    config.experience_json = body.experience
    config.safety_json = body.safety
    agent.system_prompt = body.prompt_template or agent.system_prompt
    await audit(db, ctx, "agent.config.update", "agent", agent.id, {"agent_type": body.agent_type})
    await db.commit()
    return {"agent_id": agent.id, "status": "draft", "updated_at": config.updated_at}


@app.post("/api/agents/{agent_id}/publish")
async def publish_agent(agent_id: str, body: AgentPublishRequest, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    agent = await db.scalar(select(Agent).where(Agent.id == agent_id, Agent.tenant_id == ctx.tenant_id))
    if not agent:
        raise HTTPException(status_code=404, detail="智能体不存在")
    config = await get_or_create_agent_config(db, ctx, agent)
    version = body.version or (agent.version + 1)
    if await db.scalar(select(AgentVersion).where(AgentVersion.agent_id == agent.id, AgentVersion.version == version)):
        raise HTTPException(status_code=409, detail="智能体版本已存在")
    snapshot = {"agent_type": config.agent_type, "prompt_mode": config.prompt_mode, "prompt_template": config.prompt_template, "variables": config.variables_json, "memory": config.memory_json, "capabilities": config.capabilities_json, "experience": config.experience_json, "safety": config.safety_json, "provider_id": agent.provider_id, "knowledge_base_id": agent.knowledge_base_id}
    db.add(AgentVersion(tenant_id=ctx.tenant_id, agent_id=agent.id, version=version, config_json=snapshot, changelog=body.changelog, status="published", created_by=ctx.user_id))
    agent.version = version
    config.published_version = version
    config.publish_status = "published"
    await audit(db, ctx, "agent.publish", "agent", agent.id, {"version": version, "changelog": body.changelog})
    await db.commit()
    return {"agent_id": agent.id, "version": version, "status": "published"}


@app.get("/api/agents/{agent_id}/stats")
async def agent_stats(agent_id: str, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    agent = await db.scalar(select(Agent).where(Agent.id == agent_id, Agent.tenant_id == ctx.tenant_id))
    if not agent:
        raise HTTPException(status_code=404, detail="智能体不存在")
    messages = await db.scalar(select(func.count()).select_from(AgentEvent).where(AgentEvent.tenant_id == ctx.tenant_id, AgentEvent.agent_id == agent_id, AgentEvent.event_type == "message")) or 0
    tokens = await db.scalar(select(func.coalesce(func.sum(AgentEvent.tokens), 0)).where(AgentEvent.tenant_id == ctx.tenant_id, AgentEvent.agent_id == agent_id)) or 0
    active_users = await db.scalar(select(func.count(func.distinct(AgentEvent.user_id))).where(AgentEvent.tenant_id == ctx.tenant_id, AgentEvent.agent_id == agent_id, AgentEvent.event_type == "message")) or 0
    avg_latency = await db.scalar(select(func.coalesce(func.avg(AgentEvent.latency_ms), 0)).where(AgentEvent.tenant_id == ctx.tenant_id, AgentEvent.agent_id == agent_id)) or 0
    positive = await db.scalar(select(func.count()).select_from(AgentEvent).where(AgentEvent.tenant_id == ctx.tenant_id, AgentEvent.agent_id == agent_id, AgentEvent.feedback == "positive")) or 0
    feedback_total = await db.scalar(select(func.count()).select_from(AgentEvent).where(AgentEvent.tenant_id == ctx.tenant_id, AgentEvent.agent_id == agent_id, AgentEvent.feedback.in_(["positive", "negative"]))) or 0
    blocked = await db.scalar(select(func.count()).select_from(AuditEvent).where(AuditEvent.tenant_id == ctx.tenant_id, AuditEvent.resource_type == "agent", AuditEvent.resource_id == agent_id, AuditEvent.action == "safety.block")) or 0
    series_rows = (await db.execute(select(func.date_trunc("day", AgentEvent.created_at).label("day"), func.count().label("messages"), func.coalesce(func.sum(AgentEvent.tokens), 0).label("tokens")).where(AgentEvent.tenant_id == ctx.tenant_id, AgentEvent.agent_id == agent_id).group_by("day").order_by("day").limit(90))).all()
    series = [{"date": row.day, "messages": int(row.messages), "tokens": int(row.tokens)} for row in series_rows]
    return {"agent_id": agent_id, "messages": int(messages), "users": int(active_users), "tokens": int(tokens), "cost": 0, "safety_blocks": int(blocked), "satisfaction": round(int(positive) / int(feedback_total) * 100, 2) if feedback_total else None, "avg_latency_ms": round(float(avg_latency), 2), "series": series}


@app.get("/api/model-catalog")
async def model_catalog(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(ModelProvider).where(ModelProvider.tenant_id == ctx.tenant_id).order_by(ModelProvider.name))).scalars().all()
    result = []
    for row in rows:
        config = row.config_json or {}
        result.append({"id": row.id, "name": row.default_model, "provider": row.name, "provider_type": row.provider_type, "model_type": config.get("model_type", "llm"), "capabilities": config.get("capabilities", ["chat", "function_call"]), "context_length": config.get("context_length", 128000), "input_price": config.get("input_price", 0), "output_price": config.get("output_price", 0), "enabled": row.enabled})
    return result


@app.get("/api/model-catalog/{provider_id}")
async def model_catalog_detail(provider_id: str, ctx: Context, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(ModelProvider).where(ModelProvider.id == provider_id, ModelProvider.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="模型不存在")
    config = row.config_json or {}
    health = await db.scalar(select(ModelHealthCheck).where(ModelHealthCheck.provider_id == row.id, ModelHealthCheck.tenant_id == ctx.tenant_id).order_by(desc(ModelHealthCheck.created_at)))
    return {"id": row.id, "name": row.default_model, "provider": row.name, "provider_type": row.provider_type, "base_url": row.base_url, "model_type": config.get("model_type", "llm"), "capabilities": config.get("capabilities", ["chat", "function_call"]), "context_length": config.get("context_length", 128000), "input_price": config.get("input_price", 0), "output_price": config.get("output_price", 0), "parameters": {"temperature": config.get("temperature", 0.7), "top_p": config.get("top_p", 0.9), "max_tokens": config.get("max_tokens", 8192)}, "health": {"ok": health.ok, "latency_ms": health.latency_ms, "detail": health.detail, "created_at": health.created_at} if health else None, "enabled": row.enabled, "config": config}


@app.post("/api/model-experience/chat")
async def model_experience_chat(body: ModelExperienceRequest, ctx: OperatorContext, db: AsyncSession = Depends(get_db)):
    stmt = select(ModelProvider).where(ModelProvider.tenant_id == ctx.tenant_id, ModelProvider.enabled.is_(True))
    if body.provider_id:
        stmt = stmt.where(ModelProvider.id == body.provider_id)
    provider = (await db.execute(stmt.order_by(ModelProvider.created_at))).scalars().first()
    if not provider:
        raise HTTPException(status_code=422, detail="尚未配置可用模型")
    result = await call_chat(db, provider, [{"role": "user", "content": body.message}], tenant_id=ctx.tenant_id)
    await audit(db, ctx, "model.experience.chat", "model_provider", provider.id, {"model": provider.default_model, "usage": result.get("usage", {})})
    await db.commit()
    return {"provider_id": provider.id, "model": provider.default_model, "text": result.get("text", ""), "usage": result.get("usage", {})}


@app.get("/api/prompt-templates")
async def list_prompt_templates(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(PromptTemplate).where(PromptTemplate.tenant_id == ctx.tenant_id).order_by(desc(PromptTemplate.updated_at)))).scalars().all()
    return [{"id": row.id, "name": row.name, "slug": row.slug, "description": row.description, "content": row.content, "variables": row.variables_json, "version": row.version, "status": row.status, "workspace_id": row.workspace_id, "created_at": row.created_at, "updated_at": row.updated_at} for row in rows]


@app.post("/api/prompt-templates")
async def create_prompt_template(body: PromptTemplateCreate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    row = PromptTemplate(tenant_id=ctx.tenant_id, workspace_id=body.workspace_id, name=body.name, slug=body.slug, description=body.description, content=body.content, variables_json=body.variables, created_by=ctx.user_id)
    db.add(row)
    await db.flush()
    await audit(db, ctx, "prompt_template.create", "prompt_template", row.id)
    await db.commit()
    return {"id": row.id, "name": row.name, "slug": row.slug, "version": row.version, "status": row.status}


@app.patch("/api/prompt-templates/{template_id}")
async def update_prompt_template(template_id: str, body: PromptTemplateCreate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(PromptTemplate).where(PromptTemplate.id == template_id, PromptTemplate.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="提示词模板不存在")
    row.name, row.description, row.content, row.variables_json = body.name, body.description, body.content, body.variables
    row.version += 1
    await audit(db, ctx, "prompt_template.update", "prompt_template", row.id, {"version": row.version})
    await db.commit()
    return {"id": row.id, "version": row.version, "status": row.status}


@app.post("/api/prompt-templates/{template_id}/render")
async def render_prompt_template(template_id: str, body: dict[str, Any], ctx: Context, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(PromptTemplate).where(PromptTemplate.id == template_id, PromptTemplate.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="提示词模板不存在")
    values = body.get("variables") if isinstance(body.get("variables"), dict) else body
    rendered = row.content
    for key, value in (values or {}).items():
        rendered = rendered.replace("{{" + str(key) + "}}", str(value))
    await audit(db, ctx, "prompt_template.render", "prompt_template", row.id, {"variables": list((values or {}).keys())})
    await db.commit()
    return {"id": row.id, "version": row.version, "content": rendered, "unresolved": sorted({part.split("}}", 1)[0].strip() for part in rendered.split("{{")[1:] if "}}" in part})}


@app.post("/api/prompt-templates/{template_id}/publish")
async def publish_prompt_template(template_id: str, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(PromptTemplate).where(PromptTemplate.id == template_id, PromptTemplate.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="提示词模板不存在")
    row.status = "published"
    await audit(db, ctx, "prompt_template.publish", "prompt_template", row.id, {"version": row.version})
    await db.commit()
    return {"id": row.id, "version": row.version, "status": row.status}


@app.post("/api/prompt-templates/{template_id}/copy")
async def copy_prompt_template(template_id: str, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    source = await db.scalar(select(PromptTemplate).where(PromptTemplate.id == template_id, PromptTemplate.tenant_id == ctx.tenant_id))
    if not source:
        raise HTTPException(status_code=404, detail="提示词模板不存在")
    row = PromptTemplate(tenant_id=ctx.tenant_id, workspace_id=source.workspace_id, name=f"{source.name} 副本", slug=f"{source.slug}-copy-{secrets.token_hex(3)}", description=source.description, content=source.content, variables_json=source.variables_json, created_by=ctx.user_id)
    db.add(row)
    await db.flush()
    await audit(db, ctx, "prompt_template.copy", "prompt_template", row.id, {"source_id": source.id})
    await db.commit()
    return {"id": row.id, "name": row.name, "slug": row.slug, "version": row.version, "status": row.status}


@app.get("/api/databases")
async def list_databases(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(DatabaseConnection).where(DatabaseConnection.tenant_id == ctx.tenant_id).order_by(DatabaseConnection.name))).scalars().all()
    return [{"id": row.id, "name": row.name, "db_type": row.db_type, "host": row.host, "port": row.port, "database_name": row.database_name, "username": row.username, "status": row.status, "schema": row.schema_json, "workspace_id": row.workspace_id, "created_at": row.created_at} for row in rows]


async def open_database_connection(row: DatabaseConnection):
    password = decrypt_secret(row.password_encrypted)
    if row.db_type == "postgresql":
        return await asyncpg.connect(host=row.host, port=row.port, user=row.username, password=password, database=row.database_name, timeout=8, command_timeout=8)
    if row.db_type == "mysql":
        return await asyncmy.connect(host=row.host, port=row.port, user=row.username, password=password, db=row.database_name, connect_timeout=8, autocommit=True)
    raise RuntimeError(f"未安装或未支持数据库驱动: {row.db_type}")


@app.post("/api/databases")
async def create_database(body: DatabaseCreate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = DatabaseConnection(tenant_id=ctx.tenant_id, workspace_id=body.workspace_id, name=body.name, db_type=body.db_type, host=body.host, port=body.port, database_name=body.database_name, username=body.username, password_encrypted=encrypt_secret(body.password), status="configured", created_by=ctx.user_id)
    db.add(row)
    await db.flush()
    await audit(db, ctx, "database.create", "database", row.id, {"db_type": row.db_type})
    await db.commit()
    return {"id": row.id, "name": row.name, "status": row.status}


@app.post("/api/databases/{database_id}/test")
async def test_database(database_id: str, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(DatabaseConnection).where(DatabaseConnection.id == database_id, DatabaseConnection.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="数据库连接不存在")
    row.status = "incomplete"
    detail = "请补充数据库主机"
    if row.host and row.db_type in {"postgresql", "mysql"}:
        try:
            connection = await open_database_connection(row)
            await connection.close()
            row.status, detail = "healthy", f"{row.db_type} 连接成功"
        except Exception as exc:
            row.status, detail = "unhealthy", f"连接失败：{type(exc).__name__}"
    elif row.host:
        row.status, detail = "configured", f"已登记 {row.db_type} 连接；当前镜像未安装该数据库驱动"
    await audit(db, ctx, "database.test", "database", row.id, {"status": row.status})
    await db.commit()
    return {"id": row.id, "status": row.status, "detail": detail}


@app.get("/api/databases/{database_id}")
async def get_database(database_id: str, ctx: Context, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(DatabaseConnection).where(DatabaseConnection.id == database_id, DatabaseConnection.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="数据库连接不存在")
    tables = (row.schema_json or {}).get("tables", [])
    return {"id": row.id, "name": row.name, "db_type": row.db_type, "host": row.host, "port": row.port, "database_name": row.database_name, "username": row.username, "status": row.status, "tables": tables, "table_count": len(tables), "synced_at": (row.schema_json or {}).get("synced_at"), "created_at": row.created_at}


@app.post("/api/databases/{database_id}/sync")
async def sync_database_schema(database_id: str, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(DatabaseConnection).where(DatabaseConnection.id == database_id, DatabaseConnection.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="数据库连接不存在")
    if row.db_type not in {"postgresql", "mysql"}:
        raise HTTPException(status_code=422, detail="当前版本支持 PostgreSQL 和 MySQL schema 同步")
    if not row.host:
        raise HTTPException(status_code=422, detail="请先配置数据库主机")
    connection = None
    try:
        connection = await open_database_connection(row)
        tables = []
        if row.db_type == "postgresql":
            table_rows = await connection.fetch("SELECT table_name, table_type FROM information_schema.tables WHERE table_schema = 'public' ORDER BY table_name")
            for table in table_rows[:500]:
                column_rows = await connection.fetch("SELECT column_name, data_type, is_nullable FROM information_schema.columns WHERE table_schema = 'public' AND table_name = $1 ORDER BY ordinal_position", table["table_name"])
                tables.append({"name": table["table_name"], "type": table["table_type"], "columns": [{"name": item["column_name"], "type": item["data_type"], "nullable": item["is_nullable"] == "YES"} for item in column_rows]})
        else:
            async with connection.cursor(asyncmy.DictCursor) as cursor:
                await cursor.execute("SELECT TABLE_NAME AS table_name, TABLE_TYPE AS table_type FROM information_schema.tables WHERE table_schema = %s ORDER BY table_name", (row.database_name,))
                table_rows = await cursor.fetchall()
                for table in table_rows[:500]:
                    await cursor.execute("SELECT COLUMN_NAME AS column_name, DATA_TYPE AS data_type, IS_NULLABLE AS is_nullable FROM information_schema.columns WHERE table_schema = %s AND table_name = %s ORDER BY ordinal_position", (row.database_name, table["table_name"]))
                    column_rows = await cursor.fetchall()
                    tables.append({"name": table["table_name"], "type": table["table_type"], "columns": [{"name": item["column_name"], "type": item["data_type"], "nullable": item["is_nullable"] == "YES"} for item in column_rows]})
        row.schema_json = {"tables": tables, "synced_at": utcnow().isoformat()}
        row.status = "healthy"
    except Exception as exc:
        row.status = "unhealthy"
        raise HTTPException(status_code=502, detail=f"数据库结构同步失败：{type(exc).__name__}") from exc
    finally:
        if connection:
            await connection.close()
    await audit(db, ctx, "database.schema.sync", "database", row.id, {"tables": len(tables)})
    await db.commit()
    return {"id": row.id, "status": row.status, "table_count": len(tables), "tables": tables}


@app.post("/api/databases/{database_id}/query")
async def query_database(database_id: str, body: DatabaseQuery, ctx: OperatorContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(DatabaseConnection).where(DatabaseConnection.id == database_id, DatabaseConnection.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="数据库连接不存在")
    normalized = body.sql.strip().lower()
    if not (normalized.startswith("select") or normalized.startswith("with")):
        raise HTTPException(status_code=422, detail="数据库问答只允许只读查询")
    if row.db_type not in {"postgresql", "mysql"}:
        raise HTTPException(status_code=422, detail="当前版本的实时连接器支持 PostgreSQL 和 MySQL")
    if ";" in body.sql or "--" in body.sql or "/*" in body.sql or "*/" in body.sql:
        raise HTTPException(status_code=422, detail="数据库问答不允许多语句或 SQL 注释")
    sql = body.sql.strip()
    if " limit " not in f" {sql.lower()} ":
        sql = f"{sql.rstrip()} LIMIT 100"
    try:
        connection = await open_database_connection(row)
        if row.db_type == "postgresql":
            await connection.execute("SET statement_timeout = '8000ms'")
            records = await connection.fetch(sql)
            columns = list(records[0].keys()) if records else []
            result_rows = [{key: (value.isoformat() if isinstance(value, datetime) else value if value is None or isinstance(value, (str, int, float, bool)) else str(value)) for key, value in record.items()} for record in records]
        else:
            async with connection.cursor(asyncmy.DictCursor) as cursor:
                await cursor.execute("SET SESSION MAX_EXECUTION_TIME=8000")
                await cursor.execute(sql)
                records = await cursor.fetchall()
                columns = [item[0] for item in cursor.description] if cursor.description else []
                result_rows = [{key: (value.isoformat() if isinstance(value, datetime) else value if value is None or isinstance(value, (str, int, float, bool)) else str(value)) for key, value in record.items()} for record in records]
        await connection.close()
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"数据库查询失败：{type(exc).__name__}") from exc
    await audit(db, ctx, "database.query", "database", row.id, {"sql_length": len(body.sql), "row_count": len(result_rows)})
    await db.commit()
    return {"database_id": row.id, "status": "completed", "sql": sql, "columns": columns, "rows": result_rows, "message": f"查询完成，共 {len(result_rows)} 行"}


@app.get("/api/marketplace")
async def list_marketplace(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(MarketplaceItem).where(MarketplaceItem.tenant_id == ctx.tenant_id, MarketplaceItem.publish_status == "published").order_by(desc(MarketplaceItem.featured), desc(MarketplaceItem.uses)))).scalars().all()
    return [{"id": row.id, "resource_type": row.resource_type, "resource_id": row.resource_id, "category": row.category, "tags": row.tags_json, "favorites": row.favorites, "uses": row.uses, "featured": row.featured, "created_at": row.created_at} for row in rows]


@app.get("/api/marketplace/{item_id}")
async def marketplace_detail(item_id: str, ctx: Context, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(MarketplaceItem).where(MarketplaceItem.id == item_id, MarketplaceItem.tenant_id == ctx.tenant_id))
    if not row or row.publish_status != "published":
        raise HTTPException(status_code=404, detail="市场资源不存在")
    resource = {"id": row.resource_id}
    if row.resource_type == "agent":
        item = await db.scalar(select(Agent).where(Agent.id == row.resource_id, Agent.tenant_id == ctx.tenant_id))
        resource = {"id": row.resource_id, "name": item.name if item else "", "description": item.description if item else "", "version": item.version if item else None}
    elif row.resource_type == "skill":
        item = await db.scalar(select(Skill).where(Skill.id == row.resource_id, Skill.tenant_id == ctx.tenant_id))
        resource = {"id": row.resource_id, "name": item.name if item else "", "description": item.description if item else "", "version": item.version if item else None}
    elif row.resource_type == "plugin":
        item = await db.scalar(select(Plugin).where(Plugin.id == row.resource_id, Plugin.tenant_id == ctx.tenant_id))
        resource = {"id": row.resource_id, "name": item.name if item else "", "endpoint": item.endpoint if item else "", "kind": item.kind if item else None}
    return {"id": row.id, "resource_type": row.resource_type, "resource_id": row.resource_id, "category": row.category, "tags": row.tags_json, "favorites": row.favorites, "uses": row.uses, "featured": row.featured, "resource": resource, "created_at": row.created_at}


@app.get("/api/admin/marketplace")
async def admin_marketplace(ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(MarketplaceItem).where(MarketplaceItem.tenant_id == ctx.tenant_id).order_by(desc(MarketplaceItem.created_at)).limit(200))).scalars().all()
    return [{"id": row.id, "resource_type": row.resource_type, "resource_id": row.resource_id, "category": row.category, "tags": row.tags_json, "publish_status": row.publish_status, "favorites": row.favorites, "uses": row.uses, "featured": row.featured, "created_at": row.created_at} for row in rows]


@app.post("/api/marketplace/{item_id}/install")
async def install_marketplace_item(item_id: str, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(MarketplaceItem).where(MarketplaceItem.id == item_id, MarketplaceItem.tenant_id == ctx.tenant_id, MarketplaceItem.publish_status == "published"))
    if not row:
        raise HTTPException(status_code=404, detail="市场资源不存在")
    row.uses += 1
    await audit(db, ctx, "marketplace.install", "marketplace_item", row.id, {"resource_type": row.resource_type, "resource_id": row.resource_id})
    await db.commit()
    return {"id": row.id, "resource_type": row.resource_type, "resource_id": row.resource_id, "uses": row.uses}


@app.post("/api/marketplace")
async def create_marketplace_item(body: MarketplaceCreate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = MarketplaceItem(tenant_id=ctx.tenant_id, resource_type=body.resource_type, resource_id=body.resource_id, category=body.category, tags_json=body.tags, created_by=ctx.user_id)
    db.add(row)
    await db.flush()
    await audit(db, ctx, "marketplace.submit", "marketplace_item", row.id, {"resource_type": row.resource_type})
    await db.commit()
    return {"id": row.id, "status": row.publish_status}


@app.post("/api/marketplace/{item_id}/publish")
async def publish_marketplace_item(item_id: str, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(MarketplaceItem).where(MarketplaceItem.id == item_id, MarketplaceItem.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="市场资源不存在")
    row.publish_status = "published"
    row.published_at = utcnow()
    await audit(db, ctx, "marketplace.publish", "marketplace_item", row.id, {"resource_type": row.resource_type})
    await db.commit()
    return {"id": row.id, "status": row.publish_status}


@app.post("/api/favorites")
async def toggle_favorite(body: FavoriteRequest, ctx: Context, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(Favorite).where(Favorite.tenant_id == ctx.tenant_id, Favorite.user_id == ctx.user_id, Favorite.resource_type == body.resource_type, Favorite.resource_id == body.resource_id))
    if row:
        await db.delete(row)
        item = await db.scalar(select(MarketplaceItem).where(MarketplaceItem.tenant_id == ctx.tenant_id, MarketplaceItem.resource_type == body.resource_type, MarketplaceItem.resource_id == body.resource_id))
        if item:
            item.favorites = max(0, item.favorites - 1)
        favorited = False
    else:
        db.add(Favorite(tenant_id=ctx.tenant_id, user_id=ctx.user_id, resource_type=body.resource_type, resource_id=body.resource_id))
        item = await db.scalar(select(MarketplaceItem).where(MarketplaceItem.tenant_id == ctx.tenant_id, MarketplaceItem.resource_type == body.resource_type, MarketplaceItem.resource_id == body.resource_id))
        if item:
            item.favorites += 1
        favorited = True
    await db.commit()
    return {"favorited": favorited}


@app.get("/api/publish/channels")
async def list_publish_channels(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(PublishChannel).where(PublishChannel.tenant_id == ctx.tenant_id).order_by(PublishChannel.name))).scalars().all()
    return [{"id": row.id, "name": row.name, "code": row.code, "type": row.channel_type, "enabled": row.enabled, "approval_required": row.approval_required, "config": row.config_json} for row in rows]


@app.get("/api/publish/bindings")
async def list_publish_bindings(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(PublishBinding, PublishChannel).join(PublishChannel, PublishChannel.id == PublishBinding.channel_id).where(PublishBinding.tenant_id == ctx.tenant_id).order_by(desc(PublishBinding.created_at)).limit(200))).all()
    return [{"id": binding.id, "resource_type": binding.resource_type, "resource_id": binding.resource_id, "channel_id": binding.channel_id, "channel": channel.name, "version": binding.version, "status": binding.status, "config": binding.config_json, "published_at": binding.published_at, "created_at": binding.created_at} for binding, channel in rows]


@app.post("/api/publish/bindings")
async def create_publish_binding(body: PublishBindingCreate, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    channel = await db.scalar(select(PublishChannel).where(PublishChannel.id == body.channel_id, PublishChannel.tenant_id == ctx.tenant_id, PublishChannel.enabled.is_(True)))
    if not channel:
        raise HTTPException(status_code=404, detail="发布渠道不存在或已停用")
    row = PublishBinding(tenant_id=ctx.tenant_id, resource_type=body.resource_type, resource_id=body.resource_id, channel_id=body.channel_id, version=body.version, config_json=body.config, created_by=ctx.user_id)
    db.add(row)
    try:
        await db.flush()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(status_code=409, detail="该资源已经绑定此发布渠道") from exc
    await audit(db, ctx, "publish.binding.create", "publish_binding", row.id, {"channel": channel.code})
    await db.commit()
    return {"id": row.id, "status": row.status, "channel": channel.code}


@app.post("/api/publish/bindings/{binding_id}/publish")
async def publish_binding(binding_id: str, ctx: EditorContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(PublishBinding).where(PublishBinding.id == binding_id, PublishBinding.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="发布绑定不存在")
    row.status = "published"
    row.published_at = utcnow()
    await audit(db, ctx, "publish.binding.publish", "publish_binding", row.id, {"version": row.version})
    await db.commit()
    return {"id": row.id, "status": row.status, "published_at": row.published_at}


@app.get("/api/approvals")
async def list_approvals(ctx: Context, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(ApprovalRequestModel).where(ApprovalRequestModel.tenant_id == ctx.tenant_id).order_by(desc(ApprovalRequestModel.created_at)).limit(200))).scalars().all()
    return [{"id": row.id, "type": row.request_type, "title": row.title, "resource_type": row.resource_type, "resource_id": row.resource_id, "status": row.status, "applicant_id": row.applicant_id, "approver_id": row.approver_id, "comment": row.comment, "created_at": row.created_at, "decided_at": row.decided_at} for row in rows]


@app.post("/api/approvals")
async def create_approval(body: dict[str, Any], ctx: Context, db: AsyncSession = Depends(get_db)):
    row = ApprovalRequestModel(tenant_id=ctx.tenant_id, applicant_id=ctx.user_id, request_type=str(body.get("type") or "resource.publish"), resource_type=str(body.get("resource_type") or ""), resource_id=str(body.get("resource_id") or ""), title=str(body.get("title") or "资源审批"), payload_json=body)
    db.add(row)
    await db.flush()
    await audit(db, ctx, "approval.create", "approval_request", row.id, {"type": row.request_type})
    await db.commit()
    return {"id": row.id, "status": row.status}


@app.post("/api/approvals/{approval_id}/decision")
async def decide_approval(approval_id: str, body: ApprovalDecision, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = await db.scalar(select(ApprovalRequestModel).where(ApprovalRequestModel.id == approval_id, ApprovalRequestModel.tenant_id == ctx.tenant_id))
    if not row:
        raise HTTPException(status_code=404, detail="审批申请不存在")
    row.approver_id = ctx.user_id
    row.status = "approved" if body.approved else "rejected"
    row.comment = body.comment
    row.decided_at = utcnow()
    await audit(db, ctx, "approval.decide", "approval_request", row.id, {"approved": body.approved})
    await db.commit()
    return {"id": row.id, "status": row.status, "decided_at": row.decided_at}


@app.get("/api/intervention-rules")
async def list_intervention_rules(ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(InterventionRule).where(InterventionRule.tenant_id == ctx.tenant_id).order_by(InterventionRule.priority, InterventionRule.name))).scalars().all()
    return [{"id": row.id, "name": row.name, "match_mode": row.match_mode, "match_content": row.match_content, "scope": row.scope_json, "preset_reply": row.preset_reply, "priority": row.priority, "enabled": row.enabled, "starts_at": row.starts_at, "ends_at": row.ends_at} for row in rows]


@app.post("/api/intervention-rules")
async def create_intervention_rule(body: InterventionRuleCreate, ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    row = InterventionRule(tenant_id=ctx.tenant_id, name=body.name, match_mode=body.match_mode, match_content=body.match_content, scope_json=body.scope, preset_reply=body.preset_reply, priority=body.priority, enabled=body.enabled, created_by=ctx.user_id)
    db.add(row)
    await db.flush()
    await audit(db, ctx, "intervention_rule.create", "intervention_rule", row.id)
    await db.commit()
    return {"id": row.id, "name": row.name, "enabled": row.enabled}


@app.get("/api/admin/analytics/summary")
async def admin_analytics_summary(ctx: AdminContext, db: AsyncSession = Depends(get_db)):
    values = {}
    values["users"] = int(await db.scalar(select(func.count()).select_from(Membership).where(Membership.tenant_id == ctx.tenant_id)) or 0)
    for label, model in [("workspaces", Workspace), ("agents", Agent), ("models", ModelProvider), ("knowledge_bases", KnowledgeBase), ("workflows", Workflow), ("plugins", Plugin), ("skills", Skill)]:
        values[label] = int(await db.scalar(select(func.count()).select_from(model).where(model.tenant_id == ctx.tenant_id)) or 0)
    usage = (await db.execute(select(func.coalesce(func.sum(ComputeUsage.tokens), 0), func.coalesce(func.sum(ComputeUsage.cpu_milliseconds), 0), func.coalesce(func.sum(ComputeUsage.gpu_seconds), 0)).where(ComputeUsage.tenant_id == ctx.tenant_id))).one()
    values["usage"] = {"tokens": int(usage[0]), "cpu_milliseconds": int(usage[1]), "gpu_seconds": int(usage[2])}
    return values


@app.get("/api/admin/analytics/series")
async def admin_analytics_series(ctx: AdminContext, db: AsyncSession = Depends(get_db), days: int = 30):
    days = max(1, min(days, 180))
    since = utcnow() - timedelta(days=days - 1)
    events = (await db.execute(select(AgentEvent).where(AgentEvent.tenant_id == ctx.tenant_id, AgentEvent.created_at >= since).order_by(AgentEvent.created_at))).scalars().all()
    buckets = {}
    for offset in range(days):
        key = (since + timedelta(days=offset)).date().isoformat()
        buckets[key] = {"date": key, "messages": 0, "tokens": 0, "latency_ms": 0, "positive": 0, "negative": 0}
    for event in events:
        key = event.created_at.date().isoformat()
        bucket = buckets.get(key)
        if not bucket:
            continue
        bucket["messages"] += 1
        bucket["tokens"] += int(event.tokens or 0)
        bucket["latency_ms"] += int(event.latency_ms or 0)
        if event.feedback == "positive":
            bucket["positive"] += 1
        elif event.feedback == "negative":
            bucket["negative"] += 1
    for bucket in buckets.values():
        bucket["avg_latency_ms"] = round(bucket["latency_ms"] / bucket["messages"], 2) if bucket["messages"] else 0
        del bucket["latency_ms"]
    return {"days": days, "series": list(buckets.values())}
