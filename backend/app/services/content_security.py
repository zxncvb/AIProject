"""Enterprise upload inspection and text extraction.

The built-in checks are always available.  ClamAV and OCR are optional
enterprise integrations configured through environment variables.
"""
from __future__ import annotations

import asyncio
import io
import json
import re
import struct
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path

import httpx
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from ..config import get_settings


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}
TEXT_EXTENSIONS = {".txt", ".md", ".json", ".csv", ".log"}
OFFICE_EXTENSIONS = {".docx", ".xlsx", ".xlsm"}
DOCUMENT_EXTENSIONS = TEXT_EXTENSIONS | OFFICE_EXTENSIONS | {".pdf"} | IMAGE_EXTENSIONS


@dataclass
class InspectionResult:
    status: str = "clean"
    engine: str = "builtin"
    malware: list[str] = field(default_factory=list)
    dlp_findings: list[dict[str, str | int]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def payload(self) -> dict:
        return asdict(self)


def extension_of(filename: str) -> str:
    return Path(filename or "").suffix.lower()


def validate_file_signature(filename: str, raw: bytes) -> None:
    extension = extension_of(filename)
    if extension not in DOCUMENT_EXTENSIONS:
        raise ValueError("不支持的文件类型")
    if raw.startswith((b"MZ", b"\x7fELF")):
        raise ValueError("文件内容是可执行程序，已拒绝上传")
    if extension == ".pdf" and raw and not raw.startswith(b"%PDF-"):
        raise ValueError("PDF 文件头无效")
    if extension in OFFICE_EXTENSIONS:
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                names = {item.filename.lower() for item in archive.infolist()}
                if extension == ".docx" and "word/document.xml" not in names:
                    raise ValueError("DOCX 文件结构无效")
                if extension in {".xlsx", ".xlsm"} and "xl/workbook.xml" not in names:
                    raise ValueError("Excel 文件结构无效")
        except zipfile.BadZipFile as exc:
            raise ValueError("Office 文件结构无效") from exc


def inspect_archive(filename: str, raw: bytes) -> list[str]:
    settings = get_settings()
    warnings: list[str] = []
    if extension_of(filename) not in OFFICE_EXTENSIONS:
        return warnings
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        unpacked = sum(item.file_size for item in entries)
        if unpacked > settings.upload_max_uncompressed_bytes:
            raise ValueError("压缩文档展开后超过平台限制")
        if len(entries) > settings.upload_max_archive_entries:
            raise ValueError("压缩文档包含过多文件")
        if any("vbaproject.bin" in item.filename.lower() for item in entries):
            if not settings.upload_allow_macros:
                raise ValueError("文档包含宏代码，当前安全策略禁止上传")
            warnings.append("检测到 Office 宏代码")
    return warnings


async def _scan_clamav(raw: bytes) -> tuple[bool, str]:
    settings = get_settings()
    reader, writer = await asyncio.wait_for(
        asyncio.open_connection(settings.clamav_host, settings.clamav_port),
        timeout=settings.content_scan_timeout_seconds,
    )
    try:
        writer.write(b"zINSTREAM\0")
        for offset in range(0, len(raw), 65536):
            chunk = raw[offset : offset + 65536]
            writer.write(struct.pack(">I", len(chunk)))
            writer.write(chunk)
        writer.write(struct.pack(">I", 0))
        await writer.drain()
        response = (await asyncio.wait_for(reader.read(4096), timeout=settings.content_scan_timeout_seconds)).decode("utf-8", "replace").strip("\0\r\n ")
        return response.endswith("OK"), response
    finally:
        writer.close()
        await writer.wait_closed()


def inspect_dlp(text: str) -> list[dict[str, str | int]]:
    patterns = {
        "email": r"(?<![\w.+-])[\w.+-]+@[\w-]+(?:\.[\w-]+)+",
        "mobile_cn": r"(?<!\d)1[3-9]\d{9}(?!\d)",
        "id_card_cn": r"(?<!\d)\d{17}[\dXx](?!\d)",
        "bank_card": r"(?<!\d)\d{16,19}(?!\d)",
        "private_key": r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----",
        "access_token": r"(?i)(?:access[_-]?token|api[_-]?key|secret)\s*[:=]\s*[A-Za-z0-9_./+\-=]{12,}",
    }
    findings: list[dict[str, str | int]] = []
    for category, pattern in patterns.items():
        count = len(re.findall(pattern, text))
        if count:
            findings.append({"category": category, "count": count, "severity": "high" if category in {"private_key", "access_token"} else "medium"})
    return findings


async def inspect_upload(filename: str, raw: bytes, text: str = "") -> InspectionResult:
    settings = get_settings()
    validate_file_signature(filename, raw)
    result = InspectionResult(warnings=inspect_archive(filename, raw))
    # EICAR is intentionally detected by a fragment so the full test signature
    # never needs to be stored in this repository.
    if b"EICAR-STANDARD-ANTIVIRUS-TEST-FILE" in raw:
        result.malware.append("EICAR-Test-File")
    if settings.clamav_host:
        try:
            clean, detail = await _scan_clamav(raw)
            result.engine = "clamav+builtin"
            if not clean:
                result.malware.append(detail[:500])
        except Exception as exc:
            result.warnings.append(f"ClamAV 不可用: {type(exc).__name__}")
            if settings.clamav_required:
                raise ValueError("企业病毒扫描服务不可用，按安全策略拒绝上传") from exc
    elif settings.clamav_required:
        raise ValueError("企业病毒扫描服务尚未配置")
    if text:
        result.dlp_findings = inspect_dlp(text)
    if result.malware:
        result.status = "blocked"
    elif result.dlp_findings:
        result.status = "sensitive"
    if result.status == "blocked":
        raise ValueError(f"文件安全扫描未通过: {', '.join(result.malware)}")
    if result.dlp_findings and settings.upload_dlp_action == "block":
        categories = "、".join(str(item["category"]) for item in result.dlp_findings)
        raise ValueError(f"文件包含受策略保护的敏感内容: {categories}")
    return result


async def _ocr(filename: str, raw: bytes, content_type: str) -> str:
    settings = get_settings()
    if not settings.ocr_endpoint:
        raise ValueError("文件没有可提取文本；请配置 OCR_ENDPOINT 后处理扫描件或图片")
    headers = {"Authorization": f"Bearer {settings.ocr_api_key}"} if settings.ocr_api_key else {}
    async with httpx.AsyncClient(timeout=settings.ocr_timeout_seconds, follow_redirects=False) as client:
        response = await client.post(
            settings.ocr_endpoint,
            headers=headers,
            files={"file": (filename, raw, content_type or "application/octet-stream")},
            data={"language": settings.ocr_language},
        )
        response.raise_for_status()
        payload = response.json()
    text = payload.get("text") if isinstance(payload, dict) else None
    if not isinstance(text, str) or not text.strip():
        raise ValueError("OCR 服务没有返回可用文本")
    return text


async def extract_text(filename: str, raw: bytes, content_type: str = "") -> tuple[str, str]:
    settings = get_settings()
    extension = extension_of(filename)
    validate_file_signature(filename, raw)
    source = "parser"
    if extension == ".pdf":
        try:
            reader = PdfReader(io.BytesIO(raw))
            content = "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            raise ValueError(f"PDF 解析失败: {exc}") from exc
        if len(content.strip()) < settings.ocr_min_text_chars:
            content, source = await _ocr(filename, raw, content_type or "application/pdf"), "ocr"
    elif extension == ".docx":
        try:
            document = DocxDocument(io.BytesIO(raw))
            content = "\n".join(paragraph.text for paragraph in document.paragraphs)
        except Exception as exc:
            raise ValueError(f"DOCX 解析失败: {exc}") from exc
    elif extension in {".xlsx", ".xlsm"}:
        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            lines: list[str] = []
            extracted_chars = 0
            for sheet in workbook.worksheets:
                lines.append(f"# {sheet.title}")
                for values in sheet.iter_rows(values_only=True):
                    line = "\t".join("" if value is None else str(value) for value in values)
                    extracted_chars += len(line) + 1
                    if extracted_chars > settings.knowledge_max_extracted_chars:
                        raise ValueError("文档可提取文本超过知识库上限")
                    lines.append(line)
            content = "\n".join(lines)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"Excel 解析失败: {exc}") from exc
    elif extension in IMAGE_EXTENSIONS:
        content, source = await _ocr(filename, raw, content_type), "ocr"
    else:
        content = raw.decode("utf-8", errors="replace")
        if content.startswith("\ufeff"):
            content = content[1:]
        if extension == ".json":
            try:
                content = json.dumps(json.loads(content), ensure_ascii=False, indent=2)
            except json.JSONDecodeError as exc:
                raise ValueError("JSON 文件格式无效") from exc
    if len(content) > settings.knowledge_max_extracted_chars:
        raise ValueError("文档可提取文本超过知识库上限")
    if not content.strip():
        raise ValueError("文档没有可索引的文本内容")
    return content, source
